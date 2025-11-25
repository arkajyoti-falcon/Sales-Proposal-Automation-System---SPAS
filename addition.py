import os
import streamlit as st
import pandas as pd
import json
import re
import base64
import copy
import tempfile
from io import BytesIO
from datetime import date
from dataclasses import dataclass
from typing import Dict, List
from pathlib import Path
from collections import Counter, defaultdict

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
from groq import Groq
from dotenv import load_dotenv
import pdfplumber
import ezdxf
import requests

load_dotenv()

# ==================== GROQ CLIENT SETUP ====================
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
if not GROQ_API_KEY:
    st.sidebar.error("⚠️ GROQ_API_KEY not found in .env file!")
else:
    groq_client = Groq(api_key=GROQ_API_KEY)

# ==================== DXF COMPONENT EXTRACTION ====================

UNITS = {
    0: "Unitless", 1: "inches", 2: "feet", 3: "miles",
    4: "millimeters", 5: "centimeters", 6: "meters", 7: "kilometers",
}

def _is_noise_block(name: str) -> bool:
    """Filter out anonymous / noise blocks like *U69, *D123, etc."""
    n = name.strip()
    if re.match(r"^\*U\d+$", n, re.IGNORECASE): return True
    if re.match(r"^\*D\d+$", n, re.IGNORECASE): return True
    if n.startswith("*"): return True
    return False

def _normalize_group_name(name: str) -> str:
    """Normalize raw block name to a group name."""
    n = name.strip()
    if "|" in n: n = n.split("|")[-1]
    n = re.sub(r"[_\-]+", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    n = re.sub(r"\s*\(?\d+\)?$", "", n).strip()
    return n.lower()

def extract_dxf_components(dxf_path: Path) -> dict:
    """Extract component names + counts from DXF file."""
    doc = ezdxf.readfile(str(dxf_path))
    msp = doc.modelspace()
    hdr = doc.header
    units_code = hdr.get("$INSUNITS", None)
    try:
        units_code = int(units_code) if units_code is not None else None
    except: units_code = None
    
    extmin = hdr.get("$EXTMIN", None)
    extmax = hdr.get("$EXTMAX", None)
    raw_counts: Counter[str] = Counter()
    
    for e in msp:
        try:
            if e.dxftype() == "INSERT":
                bname = e.dxf.name
                if not _is_noise_block(bname):
                    raw_counts[bname] += 1
        except: continue
    
    group_map: dict[str, dict] = defaultdict(lambda: {"total_count": 0, "examples": Counter()})
    for raw_name, cnt in raw_counts.items():
        gname = _normalize_group_name(raw_name)
        if not gname: continue
        group_map[gname]["total_count"] += cnt
        group_map[gname]["examples"][raw_name] += cnt
    
    groups = []
    for gname, data in group_map.items():
        ex_list = [{"name": n, "count": c} for n, c in data["examples"].most_common(5)]
        groups.append({"group": gname, "total_count": int(data["total_count"]), "examples": ex_list})
    groups.sort(key=lambda x: -x["total_count"])
    
    return {
        "file": dxf_path.name,
        "units_code": units_code,
        "units_name": UNITS.get(units_code, "unknown") if units_code is not None else None,
        "extents": {"min": list(extmin) if extmin else None, "max": list(extmax) if extmax else None},
        "groups": groups,
        "raw_block_counts": {k: int(v) for k, v in raw_counts.items()},
    }

def _summarise_components_for_prompt(dxf_json: dict) -> str:
    groups = dxf_json.get("groups", [])
    if not groups: return "No component groups detected."
    lines = []
    for g in groups[:40]:
        name = g.get("group", "")
        total = g.get("total_count", 0)
        ex = g.get("examples", [])
        top_example = ex[0]["name"] if ex else ""
        lines.append(f"- {name} (count: {total}, example: {top_example})")
    return "\n".join(lines)

def _normalise_to_numbered_steps(raw_text: str) -> str:
    """Force clean 1..N numbered list from GROQ output."""
    lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]
    if len(lines) == 1:
        parts = re.split(r'(?:(?<=\.)\s+)(?=\d+\.)', lines[0])
        lines = [p.strip() for p in parts if p.strip()]
    
    steps = []
    for ln in lines:
        m = re.match(r"^(\d+)[\.\)\-]\s*(.*)$", ln)
        content = m.group(2).strip() if m else ln
        if content: steps.append(content)
    
    dedup = []
    seen = set()
    for s in steps:
        key = re.sub(r"\s+", " ", s.lower())
        if key not in seen:
            seen.add(key)
            dedup.append(s)
    
    max_steps = min(len(dedup), 9) if len(dedup) >= 5 else len(dedup)
    return "\n".join([f"{i}. {content}" for i, content in enumerate(dedup[:max_steps], start=1)])

# ==================== PROCESS FLOW GENERATION ====================

def call_groq_for_process_flow(client_name: str, project_name: str, dxf_json: dict):
    """Call GROQ to generate Process Flow from DXF components."""
    safe_dxf_json = {k: v for k, v in dxf_json.items() if k != "raw_block_counts"}
    comp_summary = _summarise_components_for_prompt(safe_dxf_json)

    system_prompt = """You are a senior solution engineer writing "Process Flow of the System" for CBS proposals.
OUTPUT FORMAT: Numbered list (5-9 steps), each: "<number>. <Short Title>: <description>"
RULES:
- Base each step on DXF component groups
- Use generic terms: infeed conveyors, cross-belt sorter, output chutes
- Include counts where useful (e.g., "58 gravity chutes")
- Follow physical flow: loading → induct → CBS → chutes/PTL
- Engineering language, not marketing
- No invented modules not in DXF"""

    user_prompt = f"""Client: {client_name}
Project: {project_name}

DXF Components:
{comp_summary}

Write "Process Flow of the System" as 5-9 numbered steps based on these components."""

    resp = groq_client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": system_prompt.strip()},
            {"role": "user", "content": user_prompt.strip()},
        ],
        temperature=0.2,
        max_tokens=900,
    )
    raw_text = resp.choices[0].message.content.strip()
    clean_steps = _normalise_to_numbered_steps(raw_text)
    return clean_steps, raw_text

# ==================== MERMAID FLOWCHART GENERATION ====================

def sanitize_mermaid_for_render(code: str) -> str:
    """Minimal sanitization for Mermaid code."""
    if not code: return code
    code = code.strip()
    if code.startswith("```"):
        lines = code.split("\n")
        if lines[0].strip().startswith("```"): lines = lines[1:]
        if lines and lines[-1].strip() == "```": lines = lines[:-1]
        code = "\n".join(lines).strip()
    if code.lower().startswith("mermaid"): code = code[7:].strip()
    if not code.lower().startswith("flowchart"): return code
    return code.replace("\r\n", "\n").replace("\r", "\n")

def generate_mermaid_png(mermaid_code: str) -> tuple:
    """Render Mermaid diagram to PNG."""
    logs = []
    mermaid_str = sanitize_mermaid_for_render(mermaid_code)
    logs.append(f"Cleaned Mermaid code:\n{mermaid_str}\n")
    
    # Strategy 1: Kroki PNG with JSON
    try:
        payload = {"diagram_source": mermaid_str, "diagram_type": "mermaid", "output_format": "png"}
        resp = requests.post("https://kroki.io/mermaid/png", json=payload, 
                           headers={"Content-Type": "application/json"}, timeout=30)
        if resp.ok and resp.content and len(resp.content) > 100:
            logs.append(f"✓ Kroki PNG returned {len(resp.content)} bytes.")
            return resp.content, "\n".join(logs)
    except Exception as e:
        logs.append(f"Kroki PNG error: {repr(e)}")
    
    # Strategy 2: Mermaid.ink
    try:
        json_payload = {"code": mermaid_str, "mermaid": {"theme": "default"}}
        b64_str = base64.b64encode(json.dumps(json_payload).encode('utf-8')).decode('ascii')
        resp = requests.get(f"https://mermaid.ink/img/{b64_str}", timeout=30)
        if resp.ok and resp.content and len(resp.content) > 100:
            logs.append(f"✓ mermaid.ink returned {len(resp.content)} bytes.")
            return resp.content, "\n".join(logs)
    except Exception as e:
        logs.append(f"mermaid.ink error: {repr(e)}")
    
    raise RuntimeError(f"Failed to render Mermaid:\n" + "\n".join(logs))

def call_groq_for_mermaid(process_flow_text: str):
    """Generate Mermaid flowchart code from process flow."""
    system_prompt = """You are a diagram expert for Mermaid v11 flowcharts.
RULES:
- Start with: flowchart TD (top-down)
- Simple IDs: A, B, C, D (no special chars)
- Square brackets for labels: A[Start]
- Short labels (2-5 words)
- Use --> for arrows
- Apply colors using classDef and :::className
- Green (#90EE90) for input, Blue (#87CEEB) for process, Yellow (#FFE97F) for sorting, 
  Orange (#FFB366) for collection, Red (#FFB3B3) for rejection
- Output ONLY mermaid code, no backticks"""

    user_prompt = f"""Convert to vertical Mermaid flowchart with colors:
{process_flow_text}

Use flowchart TD, simple node IDs (A,B,C), short labels, apply color coding with classDef."""

    resp = groq_client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": system_prompt.strip()},
            {"role": "user", "content": user_prompt.strip()},
        ],
        temperature=0.0,
        max_tokens=2000,
    )
    mermaid_code = resp.choices[0].message.content.strip()
    mermaid_code = re.sub(r"^```(?:mermaid)?\s*", "", mermaid_code, flags=re.MULTILINE)
    mermaid_code = re.sub(r"\s*```$", "", mermaid_code, flags=re.MULTILINE)
    return mermaid_code.strip()

# ==================== DOCUMENT BUILDING FUNCTIONS ====================

def build_proposed_system_description_section(doc, counter, client_name, project_name, 
                                              process_flow_text, layout_png_path):
    """Build Proposed System Description section (5.0)"""
    doc.add_page_break()
    from streamlit import add_numbered_heading, add_numbered_subheading, apply_normal_style
    
    add_numbered_heading(doc, "Proposed System Description", counter=counter)
    
    # 5.1 Objective
    add_numbered_subheading(doc, "Objective", f"{counter}.1")
    objective_text = (
        "The purpose of this proposal is to present the design, manufacturing, "
        "installation, commissioning, testing, and acceptance testing of the Cross Belt Sorter "
        f"system for sorting shipments, as per {client_name} requirements."
    )
    p = doc.add_paragraph(objective_text)
    apply_normal_style(p)
    doc.add_paragraph("")
    
    # 5.2 Summary of the System (layout PNG)
    add_numbered_subheading(doc, "Summary of the System", f"{counter}.2")
    p = doc.add_paragraph(
        "The following layout view illustrates the overall arrangement of infeed conveyors, sorter loop, "
        "and output chutes for the proposed system."
    )
    apply_normal_style(p)
    
    if layout_png_path and os.path.exists(layout_png_path):
        doc.add_paragraph("")
        doc.add_picture(layout_png_path, width=Inches(6.5))
        doc.add_paragraph("")
    else:
        p = doc.add_paragraph("The detailed layout is provided separately in the attached drawing.")
        apply_normal_style(p)
    
    # 5.3 Process Flow of the System
    add_numbered_subheading(doc, "Process Flow of the System", f"{counter}.3")
    for line in process_flow_text.splitlines():
        line = line.strip()
        if not line: continue
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    doc.add_paragraph("")
    
    # 5.4 Main Benefits
    add_numbered_subheading(doc, "Main Benefits of the Proposed Solution", f"{counter}.4")
    benefits = [
        "High operational throughput.",
        "Low occupancy of floor space in the building.",
        "Narrow discharge centers for the increased number of splits in limited space.",
        (
            "FALCON's CBS can adapt to changing business requirements by adjusting its speed "
            "to match the operational throughput requirement, thereby leading to power savings "
            "and reduced system wear & tear."
        ),
    ]
    for b in benefits:
        p = doc.add_paragraph(b, style="List Bullet")
        apply_normal_style(p)

def build_concept_description_section(doc, counter, flowchart_png_bytes, drawio_url="https://app.diagrams.net/"):
    """Build Concept Description section with Mermaid flowchart"""
    doc.add_page_break()
    from streamlit import add_numbered_heading, apply_normal_style
    
    add_numbered_heading(doc, "Concept Description", counter=counter)
    
    p = doc.add_paragraph(
        "The following flowchart illustrates the high-level process flow of the proposed system. "
        "Clicking the diagram will open draw.io in a browser for editing or further detailing."
    )
    apply_normal_style(p)
    doc.add_paragraph("")
    
    # Insert hyperlinked flowchart
    try:
        image_stream = BytesIO(flowchart_png_bytes)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        inline_shape = run.add_picture(image_stream, width=Inches(6.5))
        
        # Create external hyperlink
        rel_id = doc.part.relate_to(drawio_url, RT.HYPERLINK, is_external=True)
        inline = inline_shape._inline
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel_id)
        inline_copy = copy.deepcopy(inline)
        hyperlink.append(inline_copy)
        run._r.replace(inline, hyperlink)
    except Exception as e:
        # Fallback: insert without hyperlink
        image_stream = BytesIO(flowchart_png_bytes)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(image_stream, width=Inches(6.5))
    
    doc.add_paragraph("")

# ==================== STREAMLIT UI ====================

st.set_page_config(page_title="Falcon Proposal Generator", page_icon="📄", layout="centered")

st.markdown("""
<style>
    .main-header {
        font-size: 2.5rem;
        font-weight: 700;
        color: #1f3864;
        text-align: center;
        margin-bottom: 0.5rem;
    }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-header">Falcon Proposal Generator</div>', unsafe_allow_html=True)
st.markdown("---")

# ==================== INPUT COLLECTION ====================

with st.form("proposal_form"):
    st.subheader("📋 Basic Information")
    
    col1, col2 = st.columns(2)
    with col1:
        client_name = st.text_input("Client Name *", placeholder="Enter client name")
        project_title = st.text_input("Project Title *", placeholder="Enter project name")
    
    with col2:
        dxf_file = st.file_uploader("DXF Layout File *", type=["dxf"])
        layout_png_file = st.file_uploader("Full Solution PNG (optional)", type=["png", "jpg", "jpeg"])
    
    st.divider()
    
    st.subheader("📋 Section Configuration")
    
    col1, col2 = st.columns(2)
    with col1:
        psd_include = st.checkbox("Proposed System Description", value=True)
        cd_include = st.checkbox("Concept Description (Flowchart)", value=True)
    
    submitted = st.form_submit_button("🚀 Generate Document Sections", type="primary")

# ==================== GENERATION LOGIC ====================

if submitted:
    if not client_name or not project_title:
        st.error("❌ Please fill in Client Name and Project Title")
        st.stop()
    
    if not dxf_file:
        st.error("❌ Please upload DXF Layout File")
        st.stop()
    
    if not GROQ_API_KEY:
        st.error("❌ GROQ_API_KEY not found in environment")
        st.stop()
    
    with st.spinner("🔧 Processing DXF and generating content..."):
        try:
            # Save DXF file
            tmp_dir = Path(tempfile.mkdtemp(prefix="proposal_"))
            dxf_path = tmp_dir / dxf_file.name
            dxf_path.write_bytes(dxf_file.getvalue())
            
            # Extract DXF components
            st.info("📐 Extracting DXF components...")
            dxf_json = extract_dxf_components(dxf_path)
            
            # Generate Process Flow
            process_flow_text = None
            if psd_include:
                st.info("✍️ Generating Process Flow with AI...")
                process_flow_text, _ = call_groq_for_process_flow(
                    client_name, project_title, dxf_json
                )
                st.success("✅ Process Flow generated")
                with st.expander("📄 View Process Flow"):
                    st.text(process_flow_text)
            
            # Generate Mermaid Flowchart
            flowchart_png_bytes = None
            if cd_include and process_flow_text:
                st.info("🗺️ Generating Mermaid flowchart...")
                mermaid_code = call_groq_for_mermaid(process_flow_text)
                
                with st.expander("📄 View Mermaid Code"):
                    st.code(mermaid_code, language="mermaid")
                
                flowchart_png_bytes, render_log = generate_mermaid_png(mermaid_code)
                st.success("✅ Flowchart rendered")
                
                st.image(flowchart_png_bytes, caption="Generated Flowchart", use_container_width=True)
            
            # Handle PNG layout
            layout_png_path = None
            if layout_png_file:
                layout_png_path = tmp_dir / layout_png_file.name
                layout_png_path.write_bytes(layout_png_file.getvalue())
                layout_png_path = str(layout_png_path)
            
            # Build Document
            st.info("📄 Building DOCX document...")
            doc = Document()
            
            # Set default font
            style = doc.styles['Normal']
            style.font.name = 'Calibri (Body)'
            style.font.size = Pt(11)
            
            counter = 1
            
            # Add Proposed System Description
            if psd_include and process_flow_text:
                build_proposed_system_description_section(
                    doc, counter, client_name, project_title, 
                    process_flow_text, layout_png_path
                )
                counter += 1
            
            # Add Concept Description
            if cd_include and flowchart_png_bytes:
                build_concept_description_section(doc, counter, flowchart_png_bytes)
                counter += 1
            
            # Save document
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            st.success("🎉 Document generated successfully!")
            
            st.download_button(
                label="📥 Download Proposal Sections",
                data=buffer,
                file_name=f"Proposal_Sections_{client_name.replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
            
        except Exception as e:
            st.error(f"❌ Error: {str(e)}")
            st.exception(e)

st.markdown("---")
st.info("💡 **Tip:** Make sure GROQ_API_KEY is set in your .env file")