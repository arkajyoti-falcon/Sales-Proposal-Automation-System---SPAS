import streamlit as st
from io import BytesIO
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="Principal of Safety Builder", page_icon="🦺")

st.title("Principal of Safety – Section Builder")
st.write("This will generate a fixed 'Principal of Safety' section with images into a DOCX file.")


# Helper to insert an image safely (fallback to placeholder text)
def add_center_image(doc, path, width_in=4.5):
    p = doc.add_paragraph()
    try:
        run = p.add_run()
        run.add_picture(path, width=Inches(width_in))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        p.add_run(f"[Image missing: {path}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_cell_image(cell, path, width_in=2.5):
    p = cell.add_paragraph()
    try:
        run = p.add_run()
        run.add_picture(path, width=Inches(width_in))
    except Exception:
        p.add_run(f"[Image missing: {path}]")


def build_principal_of_safety_docx() -> bytes:
    doc = Document()

    # Main heading
    doc.add_heading("Principal of Safety", level=2)

    # ------------------------------------------------
    # 1. E-Stops
    # ------------------------------------------------
    p = doc.add_paragraph()
    p.add_run("1. E-Stops").bold = True

    # image1 – center
    add_center_image(doc, "FIXED_IMAGE\\e-stop.png", width_in=4.5)

    # a. At Every Conveyor Module – Both sides
    p = doc.add_paragraph()
    p.add_run("            a. At Every Conveyor Module – Both sides").bold = False
    add_center_image(doc, "FIXED_IMAGE\\at-every.png", width_in=4.5)

    # b. VDS Chutes
    p = doc.add_paragraph()
    p.add_run("             b. VDS Chutes").bold = False
    add_center_image(doc, "FIXED_IMAGE\\emergency-vds.png", width_in=4.5)

    doc.add_paragraph("")

    # ------------------------------------------------
    # 2. Pull Cords Switch
    # ------------------------------------------------
    p = doc.add_paragraph()
    p.add_run("2. Pull Cords Switch").bold = True

    doc.add_paragraph("Required for Infeed & Takeout Conveyors")

    # image4 left, image5 right in a 2-column table
    table_pc = doc.add_table(rows=1, cols=2)
    row = table_pc.rows[0].cells
    add_cell_image(row[0], "FIXED_IMAGE\\pull-cords.PNG", width_in=3)
    add_cell_image(row[1], "FIXED_IMAGE\puul-cords-arch.PNG", width_in=3)

    doc.add_paragraph("")

    # ------------------------------------------------
    # 3. Fencing
    # ------------------------------------------------
    p = doc.add_paragraph()
    p.add_run("3. Fencing").bold = True

    # a. Between Inducts (left text, right image)
    table_fence = doc.add_table(rows=1, cols=2)
    

    # Row 0: a. Between Inducts + image6
    c00 = table_fence.rows[0].cells[0]
    c01 = table_fence.rows[0].cells[1]
    c00.text = "a. Between Inducts \nb. Between Inducts and Sorter"
    add_cell_image(c01, "FIXED_IMAGE\\fencing.png", width_in=2.5)

   

    doc.add_paragraph("")

    # ------------------------------------------------
    # 4. Leg Guards
    # ------------------------------------------------
    p = doc.add_paragraph()
    p.add_run("4. Leg Guards").bold = True

    # a. Leg Guards:
    p = doc.add_paragraph()
    p.add_run("a. Leg Guards:").bold = True
    doc.add_paragraph(
        "Leg guards are protective components designed to shield the legs or supports that "
        "stabilize machinery. They are essential for preventing accidents, injuries, and "
        "equipment damage by covering exposed areas where individuals might come into contact "
        "with moving parts or sharp edges. Material for leg guards is considered as MS."
    )

    # image7 – center
    add_center_image(doc, "FIXED_IMAGE\\leg_gurads.png", width_in=4.5)

    # b. Leg Guard Types
    p = doc.add_paragraph()
    p.add_run("b. Leg Guard Types").bold = True

    # Table with border
    table_types = doc.add_table(rows=1, cols=5)
    table_types.style = "Table Grid"
    hdr = table_types.rows[0].cells
    hdr[0].text = "Sl No."
    hdr[1].text = "Leg Guards"
    hdr[2].text = "Application"
    hdr[3].text = "Drawing"
    hdr[4].text = "Reference Image"

    # Row 1
    r1 = table_types.add_row().cells
    r1[0].text = "1"
    r1[1].text = "SPECIAL LEGS"
    r1[2].text = "Single Decker CBS Sorter & Special leg structure area for conveyors"
    add_cell_image(r1[3], "FIXED_IMAGE\\draw1.PNG", width_in=1.5)
    add_cell_image(r1[4], "FIXED_IMAGE\\ref1.PNG", width_in=1.5)

    # Row 2
    r2 = table_types.add_row().cells
    r2[0].text = "2"
    r2[1].text = "XL"
    r2[2].text = "DOUBLE DECKER / XL CBS"
    add_cell_image(r2[3], "FIXED_IMAGE\\draw2.PNG", width_in=1.5)
    add_cell_image(r2[4], "FIXED_IMAGE\\ref2.PNG", width_in=1.5)

    # Row 3
    r3 = table_types.add_row().cells
    r3[0].text = "3"
    r3[1].text = "MODULAR LEGS"
    r3[2].text = "For Modularization Leg"
    add_cell_image(r3[3], "FIXED_IMAGE\\draw3.PNG", width_in=1.5)
    add_cell_image(r3[4], "FIXED_IMAGE\\ref3.PNG", width_in=1.5)

    doc.add_paragraph("")

    # c. Leg Guard Rules
    p = doc.add_paragraph()
    p.add_run("c. Leg Guard Rules").bold = True

    # i. Single Deck/Double Deck and XL CBS...
    p = doc.add_paragraph()
    p.add_run(
        "i. Single Deck/Double Deck and XL CBS: "
    ).bold = True
    p.add_run(
        "We include leg guards on all CBS legs except in the chute areas."
    )

    add_center_image(doc, "FIXED_IMAGE\\leg_guards_rules.PNG", width_in=4.5)

    # ii. Special Leg Structure for Conveyors...
    p = doc.add_paragraph()
    p.add_run("ii. Special Leg Structure for Conveyors: ").bold = True
    p.add_run(
        "Where a 2.4m or wider aisle is required underneath the conveyor, we include leg guards for those legs."
    )

    add_center_image(doc, "FIXED_IMAGE\\leg_guards_rules2.PNG", width_in=4.5)

    # iii. Standard Legs for Conveyors & Chutes...
    p = doc.add_paragraph()
    p.add_run("iii. Standard Legs for Conveyors & Chutes: ").bold = True
    p.add_run(
        "We include leg guards for all conveyors that are running at a height greater than 1.5m."
    )

    add_center_image(doc, "FIXED_IMAGE\\leg_guards_rules3.PNG", width_in=4.5)

    doc.add_paragraph(
        "Any additional Leg guard requirement from client resulting from layout change for example "
        "conveyor location, number of pathways or conveyor heights during DAP, should be catered "
        "through Change Management process."
    )

    # Export to bytes
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# Streamlit actions
if st.button("Generate Principal of Safety DOCX"):
    doc_bytes = build_principal_of_safety_docx()
    st.session_state["principal_safety_doc"] = doc_bytes
    st.success("DOCX generated. You can download it below.")

if "principal_safety_doc" in st.session_state:
    st.download_button(
        label="Download Principal of Safety (.docx)",
        data=st.session_state["principal_safety_doc"],
        file_name="Principal_of_Safety_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
