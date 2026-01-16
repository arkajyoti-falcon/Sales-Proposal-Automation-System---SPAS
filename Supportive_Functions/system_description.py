# import streamlit as st
# import json
# from groq import Groq
# from docx import Document
# from docx.shared import Pt, RGBColor, Inches
# from docx.enum.text import WD_ALIGN_PARAGRAPH
# import io

# # Page configuration
# st.set_page_config(
#     page_title="System Description Generator",
#     page_icon="📦",
#     layout="centered"
# )

# # Initialize session state
# if 'generated_content' not in st.session_state:
#     st.session_state.generated_content = None

# ENHANCED_SYSTEM_PROMPT = """You are an expert Material Handling System Engineer specializing in Cross-Belt Sorter systems. Your task is to generate COMPREHENSIVE, DETAILED, and EXTENSIVE system descriptions that match the depth and technical detail of professional engineering documentation.

# **CRITICAL INSTRUCTIONS:**

# 1. **USE ONLY PROVIDED INFORMATION:**
#    - Extract ALL information from the process flow input
#    - Extract ALL quantities and specifications from the DXF file information
#    - DO NOT use any values from training examples
#    - DO NOT assume or invent specifications

# 2. **DXF FILE INTEGRATION:**
#    You will receive DXF file information in JSON format containing:
#    - File name and units
#    - Block counts for components (chutes, operators, leg guards, fencing, pallets, etc.)
#    - Groups with total counts
   
#    **Use this DXF data to:**
#    - Extract exact quantities for chutes, operators, safety equipment
#    - Include specific counts in relevant sections
#    - Reference the DXF file as the source of layout information
#    - Add details about protection, fencing, and infrastructure based on block counts

# 3. **SECTION GENERATION - BE EXTREMELY DETAILED:**

#    Create sections ONLY for components mentioned in process flow or DXF data. Each section must be COMPREHENSIVE with multiple paragraphs.

#    **INFEED SYSTEM** (if mentioned):
#    - Write 4-6 detailed paragraphs
#    - Describe the overall configuration and purpose
#    - Explain each conveyor type in detail (3-4 sentences each):
#      * **Straight Belt Conveyor**: Modular and robust design, used for smooth conveying of products over straight paths. MS profile is used to build conveyor frame. The conveyors are supplied with necessary supports and bolts to fix them to the supporting plane, as well as junction elements allowing easy and jam-free passage from one conveyor to another. Features include low noise, maximum uptime, minimal maintenance, high safety standards, and fastest ROI.
#      * **Inclined PVC Conveyor**: Used for smooth conveying of products over inclined and declined paths. Belt conveyors feature modular design with MS profile construction. Supplied with necessary supports, bolts, and junction elements for seamless integration.
#      * **Buffer Conveyor**: A buffer conveyor, also known as a buffering conveyor or accumulation conveyor, is a type of conveyor system used to temporarily store or hold items in a controlled manner. Its primary purpose is to manage the flow of items between different stages of a production or handling process when there is a mismatch in the speeds or capacities of the upstream and downstream equipment. These conveyors are required to maintain the throughput of the line.
#      * **Curve Conveyor**: Robust and easily maintainable design. The uniquely designed curves and belts provide smooth environment to parcels for making turns. The metal frames of the belts are not deformable to prevent belt misalignment. The belt guide assembly includes removable parts to allow quick replacement in case of damage.
#    - Mention flow path from loading to induct zone
#    - Include general specifications format: Belt material (PVC), load capacity, motor type (AC Geared Motor), gear motor makes, drive makes
#    - Reference total conveyor counts if available from DXF

#    **INDUCTION/FEEDLINE SYSTEM** (if mentioned):
#    - Write 5-8 detailed paragraphs
#    - Describe overall feedline configuration
#    - Detail each module type with 3-4 sentences:
#      * **Loading/Receiving Conveyor**: A receiving conveyor is a type of conveyor system used to receive and release the products for induction onto CBS. It serves as the connection point at turn point of entry where products are collected and conveyed to subsequent stages of the process. The receiving conveyor accurately positions parcels for smooth transfer to the main sorter.
#      * **Weighing Conveyor**: A weighing conveyor, also known as a weigh belt conveyor, is a type of conveyor system specifically designed to measure the weight of materials as they move along the conveyor belt. It combines the functions of conveying and weighing into a single integrated process. Weighing conveyors are equipped with high precision load cells to capture the weight of shipments. Makes include Bizerba, Mettler Toledo, or equivalent manufacturers.
#      * **Spacing Conveyor**: A spacing conveyor, also referred to as a gapping conveyor or gap optimizer, is a type of conveyor system used to create and maintain consistent gaps or spacing between items as they move along the conveyor line. Its primary purpose is to regulate the flow and spacing of products to ensure smooth operation and efficient downstream processes. This conveyor is a variable speed special purpose module that creates space between parcels as well as regulates feeding to downstream equipment.
#      * **Buffer Conveyors**: Used to temporarily store or hold items in controlled manner. Primary purpose is to manage flow of items between different stages when there is mismatch in speeds or capacities of upstream and downstream equipment. Required to maintain the throughput of line.
#      * **Angle Merge Conveyor**: An angle/intelligent merge conveyor incorporates advanced automation and control technologies to intelligently merge stream of materials into a single unified flow. It optimizes the merging process by dynamically adjusting the speed and position of items to ensure a smooth and efficient merge. This is typically a 30° triangular high-speed conveyor used for inducting shipments/boxes directly onto the sorter. The belts are strip belts for smooth shipment movement.
#    - Explain sensor placement and functionality
#    - Describe how parcels are prepared and positioned for sorter entry
#    - Include number of feedlines and capacity from process flow

#    **MANUAL INDUCT STATIONS** (if mentioned):
#    - Write 2-3 paragraphs
#    - Describe location (ground level, mezzanine)
#    - Explain operator workflow in detail
#    - Mention capacity and number of stations
#    - Include operator count from DXF data if available

#    **CROSS-BELT SORTER (Main Sorter)**:
#    - Write 4-6 detailed paragraphs
#    - Describe sorter type (Linear CBS or Loop CBS)
#    - Installation details: height from ground, location
#    - Carrier specifications: type (single/dual belt), pitch, belt dimensions
#    - For Linear: top running length, total length, number of carriers
#    - For Loop: loop circumference, deck configuration
#    - Operation description: How parcels pass through the sorter, barcode scanning process, chute assignment logic, carrier actuation mechanism, discharge process
#    - Explain the sorting sequence step by step

#    **BARCODE SCANNING & DIMENSIONING SYSTEM** (if mentioned):
#    - Write 3-4 paragraphs
#    - Scanner type and configuration (5-side, 6-side, top-only)
#    - Technology: ICR (Image Code Reader) or other
#    - Manufacturer and model information
#    - Capabilities: Barcode types (1D, 2D), scanning coverage, orientation
#    - Additional features: Image archiving, dimension measurement accuracy
#    - Integration with WCS and sorting logic

#    **OUTPUT CHUTES** - BE VERY DETAILED:
#    - **Use exact quantities from DXF data**
#    - Write 8-12 paragraphs total covering all chute types
   
#    For each chute type present:
   
#    **Collection Chutes / Manual Chutes**:
#    - Extract total count from DXF data (look for "chute", "Chute" in block counts)
#    - Write 3-4 paragraphs describing:
#      * Type: Friction roller chute or gravity chute design
#      * Purpose: A friction roller chute is a type of chute used for the smooth descent of materials or objects from an elevated position to a lower level. It utilizes its roller platform to gradually descend and collect the parcel at the end.
#      * Configuration: Single deck or double deck
#      * Capacity calculation with example dimensions
#      * Equipment per chute: Chute full sensors (quantity and function), three-color tower lights/beacon lights (to indicate chute status), push buttons (to start/stop sorting operations)
   
#    **Live Chutes / Live Dock Chutes** (if mentioned):
#    - Write 2-3 paragraphs
#    - Describe: A live chute refers to a combination of collection chute, PVC belt conveyor, and TBC (if applicable), where the collection chute helps bringing down the sorted parcel and releases it to running conveyor for direct loading into trucks
#    - Configuration and integration with conveyors
   
#    **Rejection/Technical Chutes** (if mentioned):
#    - Write 2-3 paragraphs
#    - Purpose: Handle rejected, oversized, overweight, no-read parcels
#    - Design and operation
#    - Equipment included
   
#    **Direct Bagging Chutes** (if applicable):
#    - Write 2-3 paragraphs
#    - Purpose and operation
#    - Integration with bagging system

#    **RECIRCULATION & MANUAL REFEED LINE** (if mentioned):
#    - Write 3-4 paragraphs
#    - Recirculation line: Strategically designed at the end of the sorter system to manage parcels that encounter sorting failures. This automated line efficiently gathers and transports the sort-failed parcels, refeeding them back into the sorter system without requiring additional manual labor. The entire process is seamless, ensuring parcels are automatically re-fed into the sorting system.
#    - Manual refeed line: Integration for reintroduction of rejected parcels that have been manually reprocessed. This ensures that manually handled parcels are easily fed back into the sorter, maintaining operational flow and minimizing delays.

#    **BAGGING SYSTEM** (if applicable):
#    - Write 3-4 paragraphs
#    - Bagging conveyor configuration
#    - Flow from bagging chutes to bag induct
#    - Bag scanning and induction process

#    **SECONDARY SORTING / PALLETIZATION** (if applicable):
#    - Write 2-3 paragraphs
#    - Operator workflow with hand-held terminals
#    - Pallet positioning and dispatch
#    - Include pallet count from DXF data if available

#    **TELESCOPIC BELT CONVEYORS** (if applicable):
#    - Write 2-3 paragraphs
#    - Quantity and placement
#    - Technical specifications: base length, extended length, belt specifications
#    - Purpose and operation

#    **INFRASTRUCTURE & SUPPORT SYSTEMS**:
#    - Write 6-10 paragraphs covering all infrastructure elements
   
#    **Mezzanine Platform** (if mentioned):
#    - Total area, clear height, type
#    - Number of staircases
#    - Deck configuration
   
#    **Safety & Protection**:
#    - **Extract counts from DXF data**:
#      * Leg guards count (look for "leg guard", "Leg Guard" in blocks)
#      * Operator safety guards (look for "operator safety" in blocks)
#      * Fencing (look for "fencing", "Fencing" in blocks)
#    - Write detailed paragraphs: Leg guards are protective components designed to shield the legs from external material or component. Material for leg guards is typically MS (Mild Steel). Operator safety guards protect personnel near the system. Perimeter fencing defines the loading zone and protects personnel.
   
#    **Pathways**:
#    - Allocated pathways for operator and vehicle movement
   
#    **System Color Coding** (if applicable):
#    - RAL color codes for different system components
   
#    **Electrical & Controls Infrastructure**:
#    - Control panels, switch racks, socket provisions
#    - Cable management systems
#    - Communication protocols

#    **SYSTEM TECHNICAL SUMMARY**:
#    - Write 3-4 paragraphs summarizing:
#      * Total conveyor system metrics
#      * Feedline configuration and capacity
#      * Sorter specifications
#      * Total chutes by type (use DXF counts)
#      * Operator positions (from DXF)
#      * Infrastructure elements
#      * Key equipment and technologies

# 4. **WRITING REQUIREMENTS:**
#    - Each major section: 4-8 paragraphs minimum
#    - Each subsection: 2-4 paragraphs minimum
#    - Each component description: 3-5 sentences minimum
#    - Use technical, professional language
#    - Explain functionality, purpose, and integration
#    - Include design rationale where applicable
#    - Maintain consistent technical depth throughout
#    - Use proper material handling terminology

# 5. **QUANTITY EXTRACTION FROM DXF:**
#    - Total chutes: Sum all chute-related blocks
#    - Operators: Look for "operator", "Operator" in block names
#    - Leg guards: Look for "leg guard", "Leg Guard"
#    - Fencing: Look for "fencing", "Fencing"
#    - Pallets: Look for "pallet", "Pallet"
#    - Safety equipment: Look for "safety", "gaurd", "guard"
#    - Use these exact numbers in relevant sections

# 6. **OUTPUT LENGTH TARGET:**
#    - Aim for 3000-5000 words total
#    - Match the depth and detail of professional engineering system descriptions
#    - Every component gets thorough explanation
#    - Multiple paragraphs per major section

# **REMEMBER:**
# - Be EXTREMELY detailed and comprehensive
# - Write multiple paragraphs for each section
# - Use exact quantities from DXF data
# - Explain every component thoroughly
# - Match the professional engineering documentation style
# - Generate content that is 5-10 pages when exported to Word"""

# def generate_system_description(process_flow: str, dxf_info: str, api_key: str, project_name: str) -> str:
#     """Generate system description using Groq API"""
#     try:
#         client = Groq(api_key=api_key)
        
#         # Parse DXF info to extract key quantities
#         dxf_data = {}
#         if dxf_info.strip():
#             try:
#                 dxf_data = json.loads(dxf_info)
#             except:
#                 pass
        
#         user_prompt = f"""Generate a COMPREHENSIVE, DETAILED system description for:

# PROJECT NAME: {project_name}

# PROCESS FLOW:
# {process_flow}

# DXF FILE INFORMATION:
# {dxf_info if dxf_info.strip() else "No DXF data provided"}

# REQUIREMENTS:
# 1. Extract ALL quantities from the DXF data (chutes, operators, leg guards, fencing, pallets)
# 2. Use these exact numbers in the appropriate sections
# 3. Generate EXTENSIVE descriptions for each component (multiple paragraphs), Add Table if needed.
# 4. Only include sections for components mentioned in process flow or present in DXF data
# 5. Write 3000-4000 words with technical depth matching professional engineering documentation
# 6. Each major section should have 3-4 paragraphs with subsections having 2-4 paragraphs
# 7. Each component description should have 3-5 sentences explaining functionality, design, and purpose

# Generate the detailed system description now."""

#         chat_completion = client.chat.completions.create(
#             messages=[
#                 {
#                     "role": "system",
#                     "content": ENHANCED_SYSTEM_PROMPT
#                 },
#                 {
#                     "role": "user",
#                     "content": user_prompt
#                 }
#             ],
#             model="llama-3.3-70b-versatile",
#             temperature=0.3,
#             max_tokens=5000,
#             top_p=0.9
#         )
        
#         return chat_completion.choices[0].message.content
    
#     except Exception as e:
#         st.error(f"Error generating system description: {str(e)}")
#         return None

# def create_docx(content: str, project_name: str) -> io.BytesIO:
#     """Create a Word document from the generated content"""
#     doc = Document()
    
#     # Set default font
#     style = doc.styles['Normal']
#     font = style.font
#     font.name = 'Calibri'
#     font.size = Pt(11)
    
#     # Add title
#     title = doc.add_heading('System Description', 0)
#     title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
#     # Add project name
#     project_para = doc.add_paragraph()
#     project_run = project_para.add_run(f'{project_name}')
#     project_run.bold = True
#     project_run.font.size = Pt(14)
#     project_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
#     # Add spacing
#     doc.add_paragraph()
    
#     # Split content into lines and process
#     lines = content.split('\n')
    
#     for line in lines:
#         line = line.strip()
#         if not line:
#             doc.add_paragraph()
#             continue
        
#         # Main section headings
#         if line.startswith('## '):
#             heading_text = line.replace('##', '').strip()
#             heading = doc.add_heading(heading_text, level=1)
#             heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)
#             heading.runs[0].font.size = Pt(14)
#             heading.runs[0].bold = True
        
#         # Sub-headings with **
#         elif line.startswith('**') and line.endswith('**'):
#             heading_text = line.replace('**', '').strip()
#             heading = doc.add_heading(heading_text, level=2)
#             heading.runs[0].font.color.rgb = RGBColor(0, 102, 204)
#             heading.runs[0].font.size = Pt(12)
#             heading.runs[0].bold = True
        
#         # Bullet points
#         elif line.startswith('- ') or line.startswith('* '):
#             bullet_text = line[2:].strip()
#             bullet_text = bullet_text.replace('**', '')
#             paragraph = doc.add_paragraph(bullet_text, style='List Bullet')
#             paragraph.runs[0].font.size = Pt(11)
        
#         # Numbered lists
#         elif len(line) > 2 and line[0].isdigit() and line[1:3] in ['. ', ') ']:
#             list_text = line[line.find(' ')+1:].strip()
#             list_text = list_text.replace('**', '')
#             paragraph = doc.add_paragraph(list_text, style='List Number')
#             paragraph.runs[0].font.size = Pt(11)
        
#         # Regular paragraphs
#         else:
#             line = line.replace('**', '')
#             paragraph = doc.add_paragraph(line)
#             paragraph.runs[0].font.size = Pt(11)
#             paragraph.runs[0].font.name = 'Calibri'
#             paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
#     # Save to BytesIO
#     doc_io = io.BytesIO()
#     doc.save(doc_io)
#     doc_io.seek(0)
    
#     return doc_io

# # Main UI
# st.title("System Description Generator")

# # Input fields
# project_name = st.text_input(
#     "Project Name",
#     placeholder="Enter project name (e.g., Warehouse Automation - 10K PPH Cross Belt Sorter)"
# )

# process_flow = st.text_area(
#     "Process Flow",
#     height=200,
#     placeholder="""Enter the process flow description:

# Example:
# 1. Infeed System: 3 infeed lines with inclined conveyors
# 2. Induction: 3 automatic feedlines, 3000 PPH capacity
# 3. Linear CBS: Barcode scanning, sorting operation
# 4. Output: 202 gravity chutes (segmented by fencing)
# 5. Recirculation: Automatic sortfail refeed"""
# )

# dxf_info = st.text_area(
#     "DXF File Information (JSON format)",
#     height=300,
#     placeholder="""{
#   "file": "project_layout.dxf",
#   "units_name": "millimeters",
#   "groups": [
#     {
#       "group": "chute",
#       "total_count": 202,
#       "examples": [{"name": "Chute", "count": 202}]
#     },
#     {
#       "group": "operator",
#       "total_count": 32,
#       "examples": [{"name": "Operator", "count": 32}]
#     },
#     {
#       "group": "leg guard",
#       "total_count": 24,
#       "examples": [{"name": "Leg Guard-01", "count": 24}]
#     }
#   ],
#   "raw_block_counts": {
#     "Chute": 202,
#     "Operator": 32,
#     "Leg Guard-01": 24,
#     "Fencing01": 6
#   }
# }"""
# )

# # API Key input
# api_key = st.text_input("Groq API Key", type="password", help="Enter your Groq API key from https://console.groq.com/keys")

# # Generate button
# if st.button("🚀 Generate System Description", type="primary", use_container_width=True):
#     if not api_key:
#         st.error("⚠️ Please enter your Groq API key")
#     elif not process_flow.strip():
#         st.error("⚠️ Please enter the process flow description")
#     elif not project_name.strip():
#         st.error("⚠️ Please enter the project name")
#     else:
#         with st.spinner("Generating comprehensive system description... This may take 30-60 seconds..."):
#             generated_content = generate_system_description(process_flow, dxf_info, api_key, project_name)
            
#             if generated_content:
#                 st.session_state.generated_content = generated_content
#                 st.success("✅ System description generated successfully!")
                
#                 # Show preview
#                 with st.expander("📄 Preview Generated Content", expanded=False):
#                     st.text_area("Preview", value=generated_content, height=300, disabled=True)
                
#                 # Download button
#                 doc_io = create_docx(generated_content, project_name)
                
#                 st.download_button(
#                     label="📥 Download as DOCX",
#                     data=doc_io,
#                     file_name=f"System_Description_{project_name.replace(' ', '_')}.docx",
#                     mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
#                     use_container_width=True
#                 )



import io
import json
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

import ezdxf
import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook

# ============================================================
# INIT
# ============================================================
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

st.set_page_config(page_title="System Description Generator", layout="wide")
st.title("System Description Generator")
st.caption(
    "DXF → Process Flow (AI) → 11. System description (AI) + Specification & Conveyor BOQ tables from Excel → DOCX"
)

UNITS = {
    0: "Unitless",
    1: "inches",
    2: "feet",
    3: "miles",
    4: "millimeters",
    5: "centimeters",
    6: "meters",
    7: "kilometers",
}


# ============================================================
# Utility helpers
# ============================================================
def _normalize_text(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def _is_noise_block(name: str) -> bool:
    n = name.strip()
    if re.match(r"^\*U\d+$", n, re.IGNORECASE):
        return True
    if re.match(r"^\*D\d+$", n, re.IGNORECASE):
        return True
    if n.startswith("*"):
        return True
    return False


def _normalize_group_name(name: str) -> str:
    n = name.strip()
    if "|" in n:
        n = n.split("|")[-1]
    n = re.sub(r"[_\-]+", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    n = re.sub(r"\s*\(?\d+\)?$", "", n).strip()
    return n.lower()


# ============================================================
# DXF → component summary
# ============================================================
def extract_dxf_components_for_process_flow(dxf_path: Path) -> dict:
    doc = ezdxf.readfile(str(dxf_path))
    msp = doc.modelspace()

    hdr = doc.header
    units_code = hdr.get("$INSUNITS", None)
    try:
        units_code = int(units_code) if units_code is not None else None
    except Exception:
        units_code = None

    extmin = hdr.get("$EXTMIN", None)
    extmax = hdr.get("$EXTMAX", None)

    raw_counts: Counter[str] = Counter()
    for e in msp:
        try:
            if e.dxftype() == "INSERT":
                name = e.dxf.name or ""
                if not name:
                    continue
                if _is_noise_block(name):
                    continue
                raw_counts[name] += 1
        except Exception:
            continue

    group_map: dict[str, dict] = defaultdict(
        lambda: {"total_count": 0, "examples": Counter()}
    )

    for raw_name, cnt in raw_counts.items():
        gname = _normalize_group_name(raw_name)
        if not gname:
            continue
        group_map[gname]["total_count"] += cnt
        group_map[gname]["examples"][raw_name] += cnt

    groups = []
    for gname, data in group_map.items():
        ex_list = [
            {"name": n, "count": c}
            for n, c in data["examples"].most_common(5)
        ]
        groups.append(
            {
                "group": gname,
                "total_count": int(data["total_count"]),
                "tokens": [],
                "examples": ex_list,
            }
        )

    groups.sort(key=lambda x: -x["total_count"])

    return {
        "file": dxf_path.name,
        "units_code": units_code,
        "units_name": UNITS.get(units_code, "unknown")
        if units_code is not None
        else None,
        "extents": {
            "min": list(extmin) if extmin is not None else None,
            "max": list(extmax) if extmax is not None else None,
        },
        "groups": groups,
    }


def _summarise_components_for_prompt(dxf_json: dict) -> str:
    groups = dxf_json.get("groups", [])
    if not groups:
        return "No component groups detected."

    lines = []
    for g in groups[:40]:
        name = g.get("group", "")
        total = g.get("total_count", 0)
        ex = g.get("examples", [])
        top_example = ex[0]["name"] if ex else ""
        lines.append(f"- {name} (count: {total}, example: {top_example})")
    return "\n".join(lines)


# ============================================================
# GROQ – Process Flow
# ============================================================
def _normalise_to_numbered_steps(raw_text: str) -> str:
    lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]

    if len(lines) == 1:
        parts = re.split(r"(?:(?<=\.)\s+)(?=\d+\.)", lines[0])
        lines = [p.strip() for p in parts if p.strip()]

    steps = []
    for ln in lines:
        m = re.match(r"^(\d+)[\.\)\-]\s*(.*)$", ln)
        content = m.group(2).strip() if m else ln
        if content:
            steps.append(content)

    dedup = []
    seen = set()
    for s in steps:
        key = re.sub(r"\s+", " ", s.lower())
        if key not in seen:
            seen.add(key)
            dedup.append(s)

    out_lines = []
    for i, content in enumerate(dedup[:12], start=1):
        out_lines.append(f"{i}. {content}")
    return "\n".join(out_lines)


def call_groq_for_process_flow(
    client_name: str, project_name: str, dxf_json: dict
) -> str:
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not set")

    comp_summary = _summarise_components_for_prompt(dxf_json)

    system_prompt = """
You are a senior sales engineer at Falcon Autotech explaining the system flow to a client.
Write the 'Process Flow of the System' as a numbered list with storytelling narrative.

## 🎯 SALES + BENEFITS FOCUS (CRITICAL)
- This is NOT a dry technical manual—it's a story about how the client's operations will IMPROVE
- Every step should subtly answer: "Why does this matter to the client?"
- Highlight speed, accuracy, efficiency, reduced errors, labor savings where relevant
- Make the client visualize their parcels flowing smoothly through the system

## 💡 WHY + WHAT (Always explain WHY, not just WHAT)
- DON'T: "Parcels are inducted onto the sorter"
- DO: "Parcels are smoothly inducted onto the sorter, ensuring zero jams and maximum throughput"
- DON'T: "Barcodes are scanned"
- DO: "Barcodes are scanned instantly, enabling precise routing to the correct destination without manual intervention"

## 🗣️ HUMAN STORYTELLING STYLE
- Write like you're walking the client through their future warehouse
- Use simple, conversational English—not technical jargon
- Transitional language: "From here...", "This feeds into...", "Your parcels then...", "Finally..."
- Sound confident and helpful, like a trusted advisor

Rules:
- Output ONLY a numbered list, 5–10 main steps.
- Each step: '<number>. <Short Title>: <description>'.
- Each step must naturally connect to and flow into the next step.
- Describe realistic material flow: infeed → distribution/buffer → induct → sorter → chutes / PTL / bagging / recirculation as indicated.
- Use component hints from the DXF summary (telescopic, infeed, VDS, auto induct, CBS, PTL, chute, bagging, etc.).
- No intro/outro text, no headings, no bullets outside the numbered steps.
"""

    user_prompt = f"""
Client: {client_name}
Project: {project_name}

DXF component summary:
{comp_summary}

Write the Process Flow of the System now.
"""

    payload = {
        "model": "llama-3.3-70b-versatile",
        "temperature": 0.2,
        "max_tokens": 800,
        "messages": [
            {"role": "system", "content": system_prompt.strip()},
            {"role": "user", "content": user_prompt.strip()},
        ],
    }

    resp = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={
            "Authorization": f"Bearer {GROQ_API_KEY}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=120,
    )
    resp.raise_for_status()
    data = resp.json()
    raw_text = data["choices"][0]["message"]["content"].strip()
    return _normalise_to_numbered_steps(raw_text)


# ============================================================
# Excel – load only visible sheets
# ============================================================
def load_visible_sheets_to_dfs(
    xlsx_bytes: bytes,
) -> tuple[dict[str, pd.DataFrame], list[str]]:
    """
    Returns:
      dfs: {sheet_name: DataFrame}
      visible_sheet_names: [sheet_name, ...]  (only sheets with sheet_state == 'visible')
    """
    stream = io.BytesIO(xlsx_bytes)
    wb = load_workbook(stream, data_only=True)
    visible_sheets = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]

    # re-seek for pandas
    stream.seek(0)
    xls = pd.ExcelFile(stream)
    dfs: dict[str, pd.DataFrame] = {}
    for sheet in visible_sheets:
        try:
            dfs[sheet] = xls.parse(sheet)
        except Exception:
            continue
    return dfs, visible_sheets


# ============================================================
# GROQ – choose sheet for Specification & BOQ
# ============================================================
def call_groq_for_sheet_selection(
    process_flow: str, sheet_names: list[str]
) -> tuple[str | None, str | None]:
    """
    Ask GROQ to pick the best sheet for:
      - Specification table (Specification/UOM/Remark)
      - Conveyor BOQ table (S No., Conveyor Type, EL_1, EL_2, Length, Width, Set, Family)

    Returns (spec_sheet, boq_sheet) – each may be None.
    """
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not set")

    system_prompt = """
You are helping configure an automated proposal generator.

You will receive:
- 'process_flow': numbered process flow steps.
- 'sheet_names': list of visible Excel sheet names in the costing workbook.

Task:
- Decide which single sheet is most likely to contain the Specification table
  (columns like 'Specification', 'UOM', 'Remark' – often in a sheet like 'Loop CBS', 'Sorter', etc.).
- Decide which single sheet is most likely to contain the Conveyor BOQ table
  (columns like 'S No.', 'Conveyor Type', 'EL_1', 'EL_2', 'Length', 'Width', 'Set', 'Family' – often 'Conveyors', 'Conveyor BOQ', etc.).

Use ONLY the provided sheet names. If unsure, set the value to null.

Return STRICT JSON like:
{
  "spec_sheet": "Loop CBS",
  "boq_sheet": "Conveyors"
}
"""

    user_payload = {
        "process_flow": process_flow,
        "sheet_names": sheet_names,
    }

    payload = {
        "model": "llama-3.3-70b-versatile",
        "temperature": 0.1,
        "max_tokens": 400,
        "messages": [
            {"role": "system", "content": system_prompt.strip()},
            {"role": "user", "content": json.dumps(user_payload, indent=2)},
        ],
    }

    resp = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={
            "Authorization": f"Bearer {GROQ_API_KEY}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=120,
    )
    resp.raise_for_status()
    data = resp.json()
    text = data["choices"][0]["message"]["content"].strip()

    # strip ```json fences if present
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
        text = re.sub(r"```$", "", text).strip()

    try:
        obj = json.loads(text)
    except Exception:
        # very defensive: try to extract first JSON object
        start = min(
            [p for p in [text.find("{"), text.find("[")] if p != -1],
            default=-1,
        )
        end = max(text.rfind("}"), text.rfind("]"))
        if start != -1 and end != -1 and end > start:
            obj = json.loads(text[start : end + 1])
        else:
            raise

    spec_sheet = obj.get("spec_sheet")
    boq_sheet = obj.get("boq_sheet")

    if spec_sheet not in sheet_names:
        spec_sheet = None
    if boq_sheet not in sheet_names:
        boq_sheet = None

    return spec_sheet, boq_sheet


# ============================================================
# Build SPEC + BOQ DataFrames from chosen sheets
# ============================================================
def _find_column_map(df: pd.DataFrame, targets: list[str]) -> dict[str, str]:
    """
    Case/space/extra-text insensitive mapping:
    returns {target_name: actual_df_column}
    """
    norm_cols = {_normalize_text(c): c for c in df.columns}
    col_map: dict[str, str] = {}

    for t in targets:
        nt = _normalize_text(t)
        for nc, col in norm_cols.items():
            if nc == nt or nc.startswith(nt) or nt in nc:
                if t not in col_map:
                    col_map[t] = col
                    break
    return col_map


def build_spec_df(df_raw: pd.DataFrame) -> pd.DataFrame:
    if df_raw is None:
        return pd.DataFrame(columns=["Specification", "UOM", "Remark"])
    df = df_raw.copy()
    df = df.dropna(how="all")

    targets = ["Specification", "UOM", "Remark"]
    col_map = _find_column_map(df, targets)

    if not col_map:
        return pd.DataFrame(columns=targets)

    cols_ordered = [col_map[t] for t in targets if t in col_map]
    df2 = df[cols_ordered].copy()

    # rename to standard names
    rename_map = {col_map[t]: t for t in col_map}
    df2 = df2.rename(columns=rename_map)

    if "Specification" in df2.columns:
        df2 = df2.dropna(subset=["Specification"])

    df2 = df2.dropna(how="all")
    return df2


def build_boq_df(df_raw: pd.DataFrame) -> pd.DataFrame:
    if df_raw is None:
        return pd.DataFrame(
            columns=[
                "S No.",
                "Conveyor Type",
                "EL_1",
                "EL_2",
                "Length",
                "Width",
                "Set",
                "Family",
            ]
        )

    df = df_raw.copy()
    df = df.dropna(how="all")

    targets = [
        "S No.",
        "Conveyor Type",
        "EL_1",
        "EL_2",
        "Length",
        "Width",
        "Set",
        "Family",
    ]
    col_map = _find_column_map(df, targets)

    if not col_map:
        return pd.DataFrame(columns=targets)

    cols_ordered = [col_map[t] for t in targets if t in col_map]
    df2 = df[cols_ordered].copy()

    rename_map = {col_map[t]: t for t in col_map}
    df2 = df2.rename(columns=rename_map)

    if "S No." in df2.columns:
        df2 = df2[~df2["S No."].isna()]

    df2 = df2.dropna(how="all")
    return df2


def df_to_row_dicts(df: pd.DataFrame, max_rows: int = 40) -> list[dict]:
    rows: list[dict] = []
    if df is None or df.empty:
        return rows

    for _, row in df.head(max_rows).iterrows():
        entry = {}
        for col in df.columns:
            val = row[col]
            entry[str(col)] = "" if pd.isna(val) else str(val)
        rows.append(entry)
    return rows


# ============================================================
# GROQ – System Description text (11.*)
# ============================================================
def call_groq_for_system_description(
    process_flow: str,
    spec_rows: list[dict],
    boq_rows: list[dict],
) -> list[dict]:
    """
    Returns a list of sections:
    [
      {"section_number": "11.1", "title": "Auto Infeed Conveyor", "paragraphs": ["...", "..."]},
      ...
    ]
    """
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not set")

    system_prompt = """
You are drafting section "11. System description" for a Cross Belt Sorter (CBS) proposal.

You will receive:
- 'process_flow': numbered steps of the Process Flow of the System.
- 'spec_table': rows from the conveyor Specification table (Specification/UOM/Remark).
- 'conveyor_boq': rows from the Conveyor BOQ (S No., Conveyor Type, EL_1, EL_2, Length, Width, Set, Family).

Write only the component-wise 'System description' – similar in style and depth to a typical Falcon proposal.
DO NOT repeat the process flow text verbatim. Use it as guidance for which modules exist and in what order.

Rules:
- Produce JSON with this shape (no markdown):

{
  "sections": [
    {
      "section_number": "11.1",
      "title": "Auto Infeed Conveyor",
      "paragraphs": ["...", "..."]
    },
    {
      "section_number": "11.2",
      "title": "Feedlines / Auto Induct",
      "paragraphs": ["...", "..."]
    }
  ]
}

- 'section_number' must start at "11.1" for the infeed / conveyor module that corresponds to the first process flow step.
- The first section 11.1 should clearly describe the Infeed / Auto Infeed Conveyor system.
  Use information from 'spec_table' where useful (belt width, MOC, roller material, etc.) in prose form.
- Subsequent sections (11.2, 11.3, ...) should cover:
  - Auto Induct / Feedlines
  - Manual Induct (if present)
  - Loop Cross Belt Sorter
  - Scanning & Dimensioning
  - Output chutes / PTL / bagging / recirculation as relevant
- Keep each section to 1–3 paragraphs of technical proposal language.
- DO NOT try to output any tables or images; only paragraphs.
- DO NOT include any keys other than: section_number, title, paragraphs.
"""

    user_payload = {
        "process_flow": process_flow,
        "spec_table": spec_rows,
        "conveyor_boq": boq_rows,
    }

    payload = {
        "model": "llama-3.3-70b-versatile",
        "temperature": 0.2,
        "max_tokens": 2500,
        "messages": [
            {"role": "system", "content": system_prompt.strip()},
            {"role": "user", "content": json.dumps(user_payload, indent=2)},
        ],
    }

    resp = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={
            "Authorization": f"Bearer {GROQ_API_KEY}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=180,
    )
    resp.raise_for_status()
    data = resp.json()
    text = data["choices"][0]["message"]["content"].strip()

    # strip ```json fences if present
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
        text = re.sub(r"```$", "", text).strip()

    try:
        obj = json.loads(text)
    except Exception:
        # fallback: extract inner JSON
        start = min(
            [p for p in [text.find("{"), text.find("[")] if p != -1],
            default=-1,
        )
        end = max(text.rfind("}"), text.rfind("]"))
        if start == -1 or end == -1 or end <= start:
            raise
        obj = json.loads(text[start : end + 1])

    if isinstance(obj, dict) and isinstance(obj.get("sections"), list):
        return obj["sections"]
    elif isinstance(obj, list):
        return obj
    else:
        raise ValueError("System Description JSON must contain a 'sections' list")


# ============================================================
# DOCX builder – 11. System description + 2 tables
# ============================================================
def _style_heading(text: str, level: int, doc: Document):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(14 if level == 1 else 12)


def _add_para(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def _add_table_from_df(doc: Document, df: pd.DataFrame):
    """
    Simple, clean Table Grid table with header row.
    """
    if df is None or df.empty:
        return

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"

    # header
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)

    # rows
    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            val = row[col]
            row_cells[i].text = "" if pd.isna(val) else str(val)


def build_system_description_doc(
    project_name: str,
    sections: list[dict],
    spec_df: pd.DataFrame | None,
    boq_df: pd.DataFrame | None,
) -> bytes:
    doc = Document()

    # Main heading (aligning with LCS style)
    _style_heading("11. System description", level=1, doc=doc)

    # Small intro line for project (optional)
    _add_para(doc, f"This section describes the major subsystems of the {project_name} solution.")

    # Ensure spec/boq are DataFrames (or None)
    if spec_df is not None and not isinstance(spec_df, pd.DataFrame):
        spec_df = None
    if boq_df is not None and not isinstance(boq_df, pd.DataFrame):
        boq_df = None

    inserted_spec = False
    inserted_boq = False

    for idx, sec in enumerate(sections):
        sec_num = sec.get("section_number") or f"11.{idx+1}"
        title = sec.get("title", "System Component")
        heading_text = f"{sec_num} {title}"
        _style_heading(heading_text, level=2, doc=doc)

        for para in sec.get("paragraphs", []):
            if para and isinstance(para, str):
                _add_para(doc, para.strip())

        # After the FIRST section (11.1) insert Specification & Conveyor BOQ
        if idx == 0:
            if spec_df is not None and not spec_df.empty:
                doc.add_paragraph("")
                _style_heading(f"{sec_num}.1 Conveyor Specification", level=3, doc=doc)
                _add_table_from_df(doc, spec_df)
                inserted_spec = True

            if boq_df is not None and not boq_df.empty:
                doc.add_paragraph("")
                _style_heading(f"{sec_num}.2 Conveyor BOQ", level=3, doc=doc)
                _add_table_from_df(doc, boq_df)
                inserted_boq = True

        doc.add_paragraph("")

    # If tables were not inserted but exist, append at end
    if not inserted_spec and spec_df is not None and not spec_df.empty:
        _style_heading("11.x Conveyor Specification", level=3, doc=doc)
        _add_table_from_df(doc, spec_df)
        doc.add_paragraph("")

    if not inserted_boq and boq_df is not None and not boq_df.empty:
        _style_heading("11.x Conveyor BOQ", level=3, doc=doc)
        _add_table_from_df(doc, boq_df)
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ============================================================
# STREAMLIT UI
# ============================================================
with st.form("sysdesc_form"):
    project_name = st.text_input(
        "Project / System Name",
        placeholder="e.g. Loop CBS 12K Sortation System",
    )
    client_name = st.text_input(
        "Client Name (for context only)",
        placeholder="e.g. LCS Italy",
        value="Client",
    )

    dxf_file = st.file_uploader("Upload DXF Layout", type=["dxf"])
    excel_file = st.file_uploader(
        "Upload Costing/BOQ Excel (only visible sheets will be used)", type=["xlsx"]
    )

    submitted = st.form_submit_button("Generate 11. System description DOCX")

if submitted:
    if not project_name:
        st.error("Project / System Name is required.")
        st.stop()
    if not dxf_file:
        st.error("DXF layout file is required.")
        st.stop()
    if not GROQ_API_KEY:
        st.error("GROQ_API_KEY is not set in environment / .env.")
        st.stop()

    tmp_dir = Path(tempfile.mkdtemp(prefix="sysdesc_"))
    dxf_path = tmp_dir / dxf_file.name
    dxf_path.write_bytes(dxf_file.getvalue())

    # 1) DXF → JSON
    st.markdown("### Step 1 – Extracting components from DXF")
    try:
        dxf_json = extract_dxf_components_for_process_flow(dxf_path)
        st.json(dxf_json)
    except Exception as e:
        st.error(f"DXF parsing failed: {e}")
        st.stop()

    # 2) Process Flow via GROQ
    st.markdown("### Step 2 – Generating Process Flow via GROQ")
    try:
        pf_text = call_groq_for_process_flow(
            client_name=client_name,
            project_name=project_name,
            dxf_json=dxf_json,
        )
        st.text_area("Process Flow (for reference, input to System Description):", pf_text, height=220)
    except Exception as e:
        st.error(f"GROQ call for Process Flow failed: {e}")
        st.stop()

    # 3) Load visible Excel sheets
    st.markdown("### Step 3 – Reading visible Excel sheets")
    dfs: dict[str, pd.DataFrame] = {}
    visible_sheet_names: list[str] = []
    if excel_file is not None:
        try:
            dfs, visible_sheet_names = load_visible_sheets_to_dfs(excel_file.getvalue())
            st.write("Visible Excel sheets:", visible_sheet_names)
        except Exception as e:
            st.error(f"Failed to read Excel: {e}")
    else:
        st.info("No Excel uploaded; System Description will not include tables.")

    # 4) Decide spec & BOQ sheet (via GROQ) and build DataFrames
    spec_df = None
    boq_df = None
    if visible_sheet_names:
        st.markdown("### Step 4 – Selecting sheets for Specification & Conveyor BOQ (via GROQ)")
        try:
            spec_sheet, boq_sheet = call_groq_for_sheet_selection(
                pf_text, visible_sheet_names
            )
            st.write("Chosen spec sheet:", spec_sheet)
            st.write("Chosen BOQ sheet:", boq_sheet)

            if spec_sheet and spec_sheet in dfs:
                spec_df = build_spec_df(dfs[spec_sheet])
                st.write("Specification table preview:")
                st.dataframe(spec_df.head(20))

            if boq_sheet and boq_sheet in dfs:
                boq_df = build_boq_df(dfs[boq_sheet])
                st.write("Conveyor BOQ table preview:")
                st.dataframe(boq_df.head(20))

        except Exception as e:
            st.error(f"GROQ sheet selection failed: {e}")
    else:
        st.info("No visible sheets detected or Excel not provided.")

    # 5) System Description text via GROQ
    st.markdown("### Step 5 – Generating 11. System description via GROQ")
    try:
        spec_rows = df_to_row_dicts(spec_df) if spec_df is not None else []
        boq_rows = df_to_row_dicts(boq_df) if boq_df is not None else []

        sections = call_groq_for_system_description(
            process_flow=pf_text,
            spec_rows=spec_rows,
            boq_rows=boq_rows,
        )
        st.json(sections)
    except Exception as e:
        st.error(f"GROQ call for System Description failed: {e}")
        st.stop()

    # 6) Build DOCX
    st.markdown("### Step 6 – Building DOCX")
    try:
        doc_bytes = build_system_description_doc(
            project_name=project_name,
            sections=sections,
            spec_df=spec_df,
            boq_df=boq_df,
        )
        st.success("DOCX generated successfully.")

        out_name = f"System_Description_{project_name.replace(' ', '_')}.docx"
        st.download_button(
            "Download 11. System description DOCX",
            data=doc_bytes,
            file_name=out_name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except Exception as e:
        st.error(f"Error while building DOCX: {e}")
