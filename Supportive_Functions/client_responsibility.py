import streamlit as st
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Client Responsibility Builder", page_icon="🤝")

st.title("Client Responsibility – Section Builder")

st.write(
    "Enter the client name. The app will generate the standard responsibility section "
    "with fixed bullets for Assembly & Commissioning, Tests, and Training."
)

# -------- Fixed templates (with {client} placeholder) --------

ASSEMBLY_TITLE_TPL = "{client} Responsibilities During the Assembly and Commissioning Phase"
ASSEMBLY_BULLETS = [
    "Provision of the site complex and office area facilities.",
    "The possibility of authorizing access to the site and the execution of the installation "
    "work up to 7 days a week and 24 hours a day if deemed necessary and if requested by FALCON.",
    "Free provision, during the installation phase, of the power supply necessary for the "
    "installation activities (estimated at 20 kW).",
    "Provision, during the commissioning phase, of the power supply necessary for the "
    "operation of the shipment sorting system free of charge at the date of FALCON need.",
    "Provision of the IT system functionality in accordance with the specification at the "
    "date of FALCON need.",
    "The customer is responsible for a safe working environment.",
    "The customer makes arrangements for the working area(s) to be protected against direct "
    "weather influences.",
    "The customer provides adequate lighting, heating, and ventilation to create a normal "
    "working environment.",
    "The cost of temporary storage that may be required (other than in the immediate vicinity "
    "of the installation site) is not included in the scope of delivery of this quotation. "
    "This also applies to temporary storage that may be required because materials are ready "
    "for delivery (in accordance with the schedule) but cannot be delivered due to hold-ups "
    "on the customer's side.",
    "The customer is responsible for the demarcation of aisles and danger zones with floor paint.",
]

TESTS_TITLE_TPL = "Responsibilities of {client} During the Tests"
TESTS_BULLETS = [
    "Provision of the test loads and barcode labels required for the tests.",
    "Provision of personnel required for test activities (loading and unloading operations).",
    "Provision of the necessary information to sort the shipments correctly.",
    "Verify with FALCON the quality and conformity of the test loads (labels, cartons).",
    "Will need to provide the necessary staff to collect information on the tests and to "
    "verify and confirm test results with FALCON.",
]

TRAINING_TITLE_TPL = "{client} Responsibilities During the Training"
TRAINING_BULLETS = [
    "Free from their usual work, the employees participate in the training for the duration "
    "of the training.",
    "Provision of a list of participants for each available training course 3 days before the "
    "start of the course.",
    "Provision of a classroom equipped with a whiteboard, video projector, projection screen, "
    "and enough space for desks or tables and chairs for the trainer and trained staff.",
    "Check the prerequisites of the people who have to follow the training, e.g. the "
    "qualifications of the technical staff.",
    "The invitation of staff to attend the training courses will be at the expense of {client}.",
]


def build_client_responsibility_docx(client: str) -> BytesIO:
    """Build DOCX with Client Responsibility section."""
    doc = Document()

    # Main heading
    doc.add_heading("Client Responsibility", level=2)

    # Assembly & Commissioning
    h_assembly = ASSEMBLY_TITLE_TPL.format(client=client)
    p = doc.add_paragraph()
    p.add_run(h_assembly).bold = True

    for b in ASSEMBLY_BULLETS:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_paragraph("")  # spacing

    # Tests
    h_tests = TESTS_TITLE_TPL.format(client=client)
    p = doc.add_paragraph()
    p.add_run(h_tests).bold = True

    for b in TESTS_BULLETS:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_paragraph("")

    # Training
    h_training = TRAINING_TITLE_TPL.format(client=client)
    p = doc.add_paragraph()
    p.add_run(h_training).bold = True

    for b in TRAINING_BULLETS:
        doc.add_paragraph(b.format(client=client), style="List Bullet")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# -------- Streamlit UI --------

client_name = st.text_input("Client name", value="Shadowfax")

if client_name.strip():
    client_clean = client_name.strip()

    # Preview text
    st.markdown("---")
    st.subheader("Preview")

    st.markdown(f"### {ASSEMBLY_TITLE_TPL.format(client=client_clean)}")
    for b in ASSEMBLY_BULLETS:
        st.markdown(f"- {b}")

    st.markdown(f"### {TESTS_TITLE_TPL.format(client=client_clean)}")
    for b in TESTS_BULLETS:
        st.markdown(f"- {b}")

    st.markdown(f"### {TRAINING_TITLE_TPL.format(client=client_clean)}")
    for b in TRAINING_BULLETS:
        st.markdown(f"- {b.format(client=client_clean)}")

    st.markdown("---")
    st.subheader("Download DOCX")

    docx_buffer = build_client_responsibility_docx(client_clean)
    st.download_button(
        label="Generate & Download Client Responsibility (.docx)",
        data=docx_buffer,
        file_name=f"{client_clean}_Client_Responsibility.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
else:
    st.info("Please enter a client name to generate the responsibility section.")
