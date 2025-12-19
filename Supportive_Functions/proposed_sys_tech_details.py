import io
import json
import os

import streamlit as st
import pandas as pd
from groq import Groq
from docx import Document


# ---------- HELPERS ----------

def df_to_compact_text(df: pd.DataFrame) -> str:
    """
    Convert Quote Master dataframe to compact text for GROQ.
    Every non-empty row becomes a pipe-separated list of "Col=Value".
    """
    lines = []
    headers = list(df.columns)

    for _, row in df.iterrows():
        if row.isna().all():
            continue

        pairs = []
        for h, v in zip(headers, row.values):
            if pd.isna(v) or h is None:
                continue
            v_str = str(v).strip()
            if not v_str:
                continue
            pairs.append(f"{h}={v_str}")

        if pairs:
            lines.append(" | ".join(pairs))

    return "\n".join(lines)


# ---- NEW, EXAMPLE-DRIVEN PROMPT ----

SYSTEM_PROMPT = """
You are a proposal BOM aggregation engine for Cross Belt Sorter projects.

INPUT
- You receive flattened rows from an Excel sheet called "Quote Master".
- Each row looks like:
  "S.No=1 | Category=Conveyor | Description=Straight PVC Conveyor BW 800 | UOM=m | Re=24 | ..."
- Columns can include S.No., Category, Sub-Category, Description, UOM, Costing Type, Cost,
  Unit Price, Sales Factor Requirement, Recommended Quantity, etc.

GOAL
- Build a SHORT "Proposed System Technical Details" section, split into 3 logical sub-sections:
  1) "Mechanical equipment"
  2) "Electrical Equipment"
  3) "Control System"
- The output must be highly aggregated, similar to Falcon proposal tables, NOT one row per part.

GENERAL RULES
- Strongly prefer 3–6 rows in Mechanical, 1 row in Electricals, 1 row in Control System.
- Group fine-grained parts into functional systems.
- DO NOT list every small item separately in the final table.
- DO NOT invent components that are not present in the sheet.

MECHANICAL EQUIPMENT – TYPICAL SYSTEMS
Create a small number of “systems” such as (examples, use only what makes sense):
- "Infeed System" or "Auto Infeed System"
- "Feedlines" or "Auto Induct Feedlines"
- "Sorter – 1 Loop Cross Belt Sorter"
- "Sorter Outputs"
- "Steel Works"
- Any other clear mechanical system suggested by the data.

Mapping guidance:
- Infeed / Auto Infeed System:
  - Straight PVC conveyors, S3 conveyors, gravity rollers, curves, merges, spacing conveyors that
    clearly belong to the main inbound line.
- Feedlines / Auto Induct Feedlines:
  - Receiving, weighing, spacing, buffer conveyors, angle merges that feed the sorter.
- Sorter:
  - Loop CBS sorter, carrier pitch, sorter length/height/speed, type of drive (LIM/LSM),
    supports, fencing, empty-carrier detection, product centring, dimension scanning, etc.
- Sorter Outputs:
  - PTL chutes, rejection chutes, spurs, pop-up sorter units, bag holding assemblies, bins, etc.
- Steel Works:
  - Platforms, mezzanines, stairs, ladders, fencing, safety guards, leg guards, end joints etc.

ELECTRICAL EQUIPMENT
- Usually a single line "Electricals – Consists of".
- Group all power/controls items:
  - main power distribution panel, MCC, main control panel, feedline control panels,
    sorter drive panels, VFD panels, network switches, field cabling, earthing, hooters,
    tower lamps, pull cords, emergency stops, IO cards, surge suppressors, harmonic filters, etc.
- Quantity: typically "1 Set".
- Value: usually just "Included" (unless there is a clear different summary).

CONTROL SYSTEM
- Usually a single line "Components – Consists of".
- Group all PLC / SCADA / IT / software items:
  - PLC based control system, SCADA, industrial switches, servers/IPC, CCTV/VMS,
    OCR/vision PCs, WCS / sorter control software, PTL controllers, licences, custom IT integration, etc.
- Quantity: typically "1 Set".
- Value: use simple summary such as "1 Nos" and "As per requirement" where appropriate.

DATA TO PRODUCE FOR EACH ROW
For each final row in the tables, you must produce:
- pos: integer, starting from 1 within each section.
- qty: short human-readable quantity, e.g. "1 Set", "14 Feedlines", "1 Sorter".
- description_lines: array of short text lines for the Description column:
    * Line 1: main system name (e.g. "Infeed System" or "Sorter").
    * Following lines: bullet-style sub points starting with "• ". These are
      the key components grouped into this system.
- value_lines: array of short text lines for the Value column. Use this for:
    * Total counts or dimensions of important components.
    * Example: "Straight PVC Conveyors: ~38 m total", "Curve Conveyors: 2 Nos".

IMPORTANT: think in terms of systems, not raw rows.

---------------- EXAMPLES (VERY IMPORTANT) ----------------

Example A – Aggregating an Infeed System

INPUT snippet (conceptual):
ROW: Category=Conveyor | Description=Straight PVC Conveyor BW 800 | UOM=m | Re=24
ROW: Category=Conveyor | Description=Straight PVC Conveyor BW 1000 | UOM=m | Re=14
ROW: Category=Conveyor | Description=Curve Conveyor 30 deg | UOM=Nos | Re=2

EXPECTED mechanical item:
{
  "pos": 1,
  "qty": "1 Infeed System",
  "description_lines": [
    "Infeed System",
    "• Straight PVC Conveyors",
    "• Curve Conveyors"
  ],
  "value_lines": [
    "Straight PVC Conveyors: ~38 m total",
    "Curve Conveyors: 2 Nos"
  ]
}

Example B – Aggregating Feedlines

INPUT snippet (conceptual):
Rows describing "Receiving Conveyor", "Weighing Conveyor", "Spacing Conveyor",
"Buffer Conveyor", "Angle merge" with various quantities.

EXPECTED mechanical item:
{
  "pos": 2,
  "qty": "14 Feedlines",
  "description_lines": [
    "Feedlines – Consists of",
    "• Receiving Conveyor",
    "• Weighing Conveyor",
    "• Spacing Conveyor",
    "• Buffer Conveyor",
    "• Angle merge"
  ],
  "value_lines": [
    "Receiving Conveyor: 1 Set",
    "Weighing Conveyor: 1 Set",
    "Spacing Conveyor: 3 Set",
    "Buffer Conveyor: 3 Set",
    "Angle merge: 2 Set"
  ]
}

Example C – Aggregating the Sorter

INPUT snippet (conceptual):
Rows describing a loop CBS sorter with height, length, speed, drive type,
carrier pitch and associated mechanical options.

EXPECTED mechanical item:
{
  "pos": 3,
  "qty": "1 Sorter",
  "description_lines": [
    "Sorter",
    "1 Loop Cross Belt Sorter",
    "• Sorter height",
    "• Sorter length",
    "• Sorter speed",
    "• Sorter drive",
    "• Carrier pitch",
    "Including:",
    "• Standard sorter supports",
    "• Product centring system",
    "• Dimension / barcode scanning system",
    "• Hooters and E-stops",
    "• Fencing"
  ],
  "value_lines": [
    "Height: approx 2900 mm",
    "Loop length: approx 150 m"
  ]
}

Example D – Electricals

INPUT snippet (conceptual):
Rows for power distribution panel, main control panel, feedline control panels,
sorter drive panels, network switches, field cabling, hooters, tower lamps, etc.

EXPECTED electrical section:
{
  "title": "Electrical Equipment",
  "items": [
    {
      "pos": 1,
      "qty": "1 Set",
      "description_lines": [
        "Electricals",
        "Consists of",
        "Main power distribution panel",
        "Main control panel",
        "Feedline control panels",
        "Sorter drive panels",
        "Network switches",
        "Field cabling"
      ],
      "value_lines": [
        "Included"
      ]
    }
  ]
}

Example E – Control System

INPUT snippet (conceptual):
Rows for Siemens PLC, SCADA, industrial switches, servers, sorter control software.

EXPECTED control section:
{
  "title": "Control System",
  "items": [
    {
      "pos": 1,
      "qty": "1 Set",
      "description_lines": [
        "Components",
        "Consists of",
        "PLC based control system with SCADA",
        "Industrial switch"
      ],
      "value_lines": [
        "1 Nos",
        "As per requirement"
      ]
    }
  ]
}

---------------- OUTPUT FORMAT ----------------

Return JSON ONLY in this schema:

{
  "sections": [
    {
      "title": "Mechanical equipment",
      "items": [
        {
          "pos": 1,
          "qty": "1 Set",
          "description_lines": ["..."],
          "value_lines": ["..."]
        }
      ]
    },
    {
      "title": "Electrical Equipment",
      "items": [ ... ]
    },
    {
      "title": "Control System",
      "items": [ ... ]
    }
  ]
}
"""


def call_groq_for_bom(api_key: str, sheet_text: str) -> dict:
    client = Groq(api_key=api_key)

    user_prompt = (
        "Below are flattened rows from the 'Quote Master' sheet.\n\n"
        "QUOTE_MASTER_ROWS:\n"
        + sheet_text
    )

    resp = client.chat.completions.create(
        model="meta-llama/llama-4-scout-17b-16e-instruct",
        temperature=0,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
    )

    return json.loads(resp.choices[0].message.content)


def add_bom_tables_to_doc(doc: Document, bom_json: dict):
    """Create the 14.x tables in the DOCX from the aggregated JSON."""
    doc.add_heading("14. Proposed System Technical Details", level=1)

    section_number = 1
    for section in bom_json.get("sections", []):
        title = section.get("title", "")
        items = section.get("items", [])
        if not items:
            continue

        # 14.1, 14.2, 14.3 headings
        doc.add_heading(f"14.{section_number} {title}", level=2)
        section_number += 1

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        hdr[0].text = "Pos."
        hdr[1].text = "Qty."
        hdr[2].text = "Description"
        hdr[3].text = "Value"

        for item in items:
            row = table.add_row().cells
            row[0].text = str(item.get("pos", ""))
            row[1].text = item.get("qty") or ""
            row[2].text = "\n".join(item.get("description_lines") or [])
            row[3].text = "\n".join(item.get("value_lines") or [])

        doc.add_paragraph()  # spacing


# ---------- STREAMLIT APP ----------

st.title("Proposed System Technical Details – Grouped BOM DOCX")

api_key = st.text_input("GROQ API key", type="password")
uploaded_file = st.file_uploader("Upload costing workbook (Excel with 'Quote Master')",
                                 type=["xlsx"])

if uploaded_file and api_key:
    try:
        xls = pd.ExcelFile(uploaded_file)
        sheet_name = None
        for s in xls.sheet_names:
            if s.lower().strip() == "quote master":
                sheet_name = s
                break

        if sheet_name is None:
            st.error("Sheet 'Quote Master' not found (case-insensitive).")
        else:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            sheet_text = df_to_compact_text(df)

            st.subheader("Flattened sheet preview sent to GROQ")
            preview = sheet_text[:2000]
            st.text(preview + ("\n...\n" if len(sheet_text) > len(preview) else ""))

            if st.button("Generate Grouped DOCX"):
                bom_json = call_groq_for_bom(api_key, sheet_text)

                doc = Document()
                add_bom_tables_to_doc(doc, bom_json)

                bio = io.BytesIO()
                doc.save(bio)
                bio.seek(0)

                st.download_button(
                    label="Download Proposed System Technical Details DOCX",
                    data=bio,
                    file_name="Proposed_System_Technical_Details_Grouped.docx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                )
    except Exception as e:
        st.error(f"Error: {e}")
