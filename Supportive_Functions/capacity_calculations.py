import os
import json
from io import BytesIO

import streamlit as st
import pandas as pd
from groq import Groq
from docx import Document
from docx.shared import Inches
from dotenv import load_dotenv
load_dotenv()
# --------------------------------------------------
# CONFIG
# --------------------------------------------------
GROQ_MODEL = "llama-3.3-70b-versatile"  # change if you use another
GROQ_API_KEY = os.getenv("GROQ_API_KEY")  # or hardcode for local testing

# --------------------------------------------------
# GROQ HELPER
# --------------------------------------------------

def build_capacity_prompt_from_excel(
    excel_bytes: bytes,
    client_name: str,
    project_name: str
) -> str:
    """
    Read the uploaded Excel (all sheets), dump them as CSV text,
    and build a very explicit extraction prompt for GROQ.
    We do NOT try to interpret any cell ourselves.
    """
    xls = pd.ExcelFile(BytesIO(excel_bytes))

    sheet_dumps = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet, header=None)
        # Keep as CSV-like text to preserve structure
        csv_text = df.to_csv(index=False, header=False)
        sheet_dumps.append(f"### Sheet: {sheet}\n{csv_text}")

    workbook_text = "\n\n".join(sheet_dumps)

    # IMPORTANT: we define a strict JSON schema and explicitly
    # tell GROQ to set fields to null if they are missing.
    prompt = f"""
You are an expert in interpreting throughput and capacity calculation Excel sheets
for parcel/shipment sortation systems (Loop CBS, Linear CBS, Cross Belt Sorters, etc.).

You are given a raw text dump of the complete Excel workbook used for capacity calculations.
Using ONLY the information present in the workbook (numbers and labels), you must extract
or compute the key capacity fields and return them as a single JSON object.

Context:
- Client: {client_name}
- Project: {project_name}

The workbook text follows after this instruction. It is a concatenation of all sheets, each
in CSV-like form.

IMPORTANT RULES:

1. **Use exact numbers from the workbook wherever a field is explicitly present.**
   - If a value is written in the sheet (e.g. "Sorter Speed 2 m/s", "Carrier per hour 6128"),
     prefer the sheet value instead of recomputing it.
2. **Only compute** a value if:
   - It is clearly implied (e.g. carrier_per_hour = speed_mps * 3600 / pitch_m) AND
   - It is NOT already available as a direct cell value.
3. If a field is not given and cannot be safely derived, set it explicitly to null.

KEY FIELDS (SEMANTICS):

- sorter_type:
    A short human-readable description like "Loop CBS", "Linear CBS", "Dual Belt Loop CBS"
    or "Cross Belt Sorter". Use what best matches the workbook text.

- sorter_speed_mps:
    Sorter speed in meters per second. If sheet says "Speed 2 m/s", set 2.0.

- pitch_m:
    Carrier pitch in meters. If sheet says "Pitch 1,175 mm", then pitch_m = 1.175.

- carriers_per_hour_cph:
    "Carrier per hour" / "Carriers/Hour" / "Carriers per hour" from the sheet.
    If not present, you may compute as:
      carriers_per_hour = speed_mps * 3600 / pitch_m
    and round to nearest integer.

- belts_per_hour_bph:
    "Belts per hour" / "Belts/Hour" from the sheet.
    If not present but the sorter is clearly Dual Belt, you may compute:
      belts_per_hour = carriers_per_hour * 2
    If single belt, belts_per_hour = carriers_per_hour.

- num_feedlines:
    Number of feedlines / inducts / infeed lines (e.g. "No of Feedlines", "No of Inducts").
    If the workbook has multiple such numbers, choose the one used in the capacity section.

- num_operators:
    Number of operators used in capacity calculations, if explicitly given
    (e.g. "No of Operators", "No of operators on manual induct station").
    If not given, set null.

- capacity_per_operator_pph:
    Capacity per operator in parcels/shipments per hour, if explicitly given
    (e.g. "Capacity per operator 1000 Shipments per hour"). If not given, set null.

- sorter_designed_capacity_A_pph:
    Sorter designed capacity on the parcel spectrum. Look for labels like:
    "Sorter Designed Capacity (A)", "Effective Designed Throughput of Sorter (A)",
    "Effective Designed TPH", or similar. Use the PPH/Shipments per hour value.

- feedline_designed_capacity_B_pph:
    Total feedline/induction capacity. Look for labels like:
    "Total Feedline designed capacity (B)", "Induction Capacity (B)",
    "Total Induction Capacity Designed", etc. Use the PPH value.

- effective_capacity_min_AB_pph:
    The effective designed capacity of the system.
    If the sheet already has "System designed throughput" or "Operational capacity",
    use that value.
    If not explicitly given, compute:
       effective_capacity_min_AB_pph = min(sorter_designed_capacity_A_pph,
                                           feedline_designed_capacity_B_pph)
    (if both are known).

- single_belt_pct and dual_belt_pct:
    Percentages of shipments handled on single and dual belts, if present
    (e.g. "Single Belts Shipments 91.36%", "Dual Belt Shipments 8.64%").
    Store them as numeric percentages (e.g. 91.36, 8.64).
    If not present, set them to null.

JSON SCHEMA (MANDATORY KEYS):

You MUST return exactly one JSON object with ALL of these keys:

{{
  "sorter_type": "Loop CBS or Linear CBS or Cross Belt Sorter etc.",
  "sorter_speed_mps": 2.0,
  "pitch_m": 1.175,
  "carriers_per_hour_cph": 0,
  "belts_per_hour_bph": 0,
  "num_feedlines": 0,
  "num_operators": null,
  "capacity_per_operator_pph": null,
  "sorter_designed_capacity_A_pph": 0,
  "feedline_designed_capacity_B_pph": 0,
  "effective_capacity_min_AB_pph": 0,
  "single_belt_pct": null,
  "dual_belt_pct": null
}}

RESPONSE FORMAT REQUIREMENTS (CRITICAL):

- Output MUST be **only** a JSON object.
- Do NOT include markdown, explanations, or any text outside the JSON.
- All numeric values must be raw numbers (no units, no commas, no % signs).
- If a value is unknown or not present, set it to null (not 0).

Below is the full workbook dump:

{workbook_text}
"""
    return prompt


def call_groq_for_capacity(prompt: str) -> dict:
    """
    Call GROQ with response_format=json_object so that we reliably get JSON.
    """
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not set in environment variables.")

    client = Groq(api_key=GROQ_API_KEY)

    chat_completion = client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[
            {
                "role": "system",
                "content": "You are a precise JSON data extractor. Always follow the schema exactly."
            },
            {
                "role": "user",
                "content": prompt
            },
        ],
        response_format={"type": "json_object"},
        temperature=0.0,
    )

    raw = chat_completion.choices[0].message.content
    return json.loads(raw)


# --------------------------------------------------
# DOCX HELPER
# --------------------------------------------------

def add_capacity_section_to_doc(
    doc: Document,
    client_name: str,
    project_name: str,
    cap: dict
) -> None:
    """
    Add 'Sorter System Capacity' section to an existing Document,
    using the extracted capacity dict.
    """
    # Heading
    doc.add_heading("Sorter System Capacity", level=2)

    intro_para = (
        f"The following table shows the throughput calculation for the sortation system "
        f"designed based on {client_name}'s {project_name} requirements."
    )
    doc.add_paragraph(intro_para)

    # Table: SPECIFICATION | VALUE
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"  # You can switch to any built-in style later

    hdr = table.rows[0].cells
    hdr[0].text = "SPECIFICATION"
    hdr[1].text = "VALUE"

    def fmt(value, suffix=""):
        if value is None or value == "":
            return "N/A"
        return f"{value}{suffix}"

    def add_row(label, value):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    # Fill rows
    add_row("Sorter Type", cap.get("sorter_type", ""))

    # Speed / pitch
    add_row("Sorter Speed", fmt(cap.get("sorter_speed_mps"), " m/s"))
    add_row("Pitch", fmt(cap.get("pitch_m"), " m"))

    # Capacity raw
    add_row("Carrier per hour", fmt(cap.get("carriers_per_hour_cph"), " CPH"))
    add_row("Belts per hour", fmt(cap.get("belts_per_hour_bph"), " BPH"))

    # Feedlines / operators
    add_row("No. of Feedlines", fmt(cap.get("num_feedlines")))
    add_row("No. of Operators", fmt(cap.get("num_operators")))
    add_row("Capacity per Operator", fmt(cap.get("capacity_per_operator_pph"), " PPH"))

    # Sorter vs Feedline capacity
    add_row(
        "Sorter Designed Capacity (A)",
        fmt(cap.get("sorter_designed_capacity_A_pph"), " PPH"),
    )
    add_row(
        "Feedline Designed Capacity (B)",
        fmt(cap.get("feedline_designed_capacity_B_pph"), " PPH"),
    )
    add_row(
        "Effective Designed Capacity (min of A & B)",
        fmt(cap.get("effective_capacity_min_AB_pph"), " PPH"),
    )

    # Optional single / dual belt %
    if cap.get("single_belt_pct") is not None or cap.get("dual_belt_pct") is not None:
        add_row(
            "Single Belt Shipments",
            fmt(cap.get("single_belt_pct"), " %"),
        )
        add_row(
            "Dual Belt Shipments",
            fmt(cap.get("dual_belt_pct"), " %"),
        )


def build_capacity_docx(
    client_name: str,
    project_name: str,
    cap: dict
) -> BytesIO:
    """
    Create a DOCX with only the Capacity section.
    You can later merge this into your main proposal builder.
    """
    doc = Document()
    add_capacity_section_to_doc(doc, client_name, project_name, cap)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# --------------------------------------------------
# STREAMLIT UI
# --------------------------------------------------

st.set_page_config(page_title="Capacity Calculations – Falcon Proposal Builder", page_icon="📈")

st.title("Capacity Calculations – DOCX Generator (via GROQ)")

st.write(
    "Upload the **throughput / capacity Excel** and this tool will ask GROQ to extract the "
    "key capacity numbers, then generate a properly formatted **Capacity Calculations** "
    "section in DOCX."
)

client_name = st.text_input("Client Name", value="Zepto")
project_name = st.text_input("Project / System Name", value="Loop CBS Sorter Project")

uploaded_excel = st.file_uploader(
    "Upload capacity Excel (Loop CBS / Linear CBS)",
    type=["xlsx", "xls"],
)

if st.button("Generate Capacity DOCX"):
    if not uploaded_excel:
        st.error("Please upload the capacity Excel file first.")
    elif not client_name.strip() or not project_name.strip():
        st.error("Please fill in Client Name and Project / System Name.")
    else:
        try:
            with st.spinner("Reading Excel and calling GROQ for capacity extraction..."):
                excel_bytes = uploaded_excel.read()
                prompt = build_capacity_prompt_from_excel(
                    excel_bytes=excel_bytes,
                    client_name=client_name.strip(),
                    project_name=project_name.strip(),
                )
                cap_data = call_groq_for_capacity(prompt)

            st.success("Capacity data extracted from GROQ:")
            st.json(cap_data)

            docx_buffer = build_capacity_docx(
                client_name=client_name.strip(),
                project_name=project_name.strip(),
                cap=cap_data,
            )

            st.download_button(
                label="Download Capacity Calculations (.docx)",
                data=docx_buffer,
                file_name=f"Capacity_Calculations_{project_name.strip().replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

        except Exception as e:
            st.error(f"Error while generating capacity section: {e}")
