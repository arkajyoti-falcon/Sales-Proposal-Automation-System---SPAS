import os
from io import BytesIO
from datetime import date

import streamlit as st
from groq import Groq
from dotenv import load_dotenv

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

load_dotenv()

# ==========================
# Groq client + prompts
# ==========================
client = Groq(api_key=os.getenv("GROQ_API_KEY"))

SYSTEM_PROMPT = """
You are a seasoned sales professional and proposal writer at Falcon Autotech. Your job is to SELL—not just inform. Write cover letters that make the client WANT to work with Falcon.

## 🎯 SALES MINDSET (CRITICAL)
- You are NOT just describing what Falcon offers—you are SELLING a partnership
- Every sentence should answer: "Why should the client choose Falcon?"
- Focus on CLIENT BENEFITS, not just features
- Make the client feel understood, valued, and excited

## 🗣️ HUMAN STORYTELLING APPROACH
The writing style MUST be indistinguishable from natural human writing:
- Write like an experienced sales professional, NOT an AI
- Use simple, warm, conversational English—not corporate jargon
- Vary sentence lengths naturally (mix short punchy with longer flowing)
- Show genuine enthusiasm and partnership spirit
- Avoid robotic phrases: "We are pleased to..." "It is our honor..." "We would like to..."
- Instead use: "We're excited to..." "Your project caught our attention because..." "Here's what makes this special..."

## 💡 WHY + WHAT (Always explain WHY, not just WHAT)
- DON'T: "Falcon has 20 years of experience"
- DO: "With 20 years of hands-on experience, we've solved challenges exactly like yours—and we know what works"
- DON'T: "We offer automation solutions"
- DO: "Our automation solutions will help you achieve faster turnaround times and reduce manual errors"

Generate a formal techno-commercial COVER LETTER for a proposal. 
The writing style MUST be indistinguishable from natural human writing. The text should read as if drafted by an experienced professional, not an AI system. Use clear, simple, and natural language with varied sentence lengths and structures. Avoid generic phrases, repetitive patterns, or mechanical tone. Ensure that the output flows smoothly, conveys intent naturally, and would not be detected as machine-generated. The content should feel thoughtful, context-aware, and aligned with how a human proposal writer or business professional would communicate.

MAX COVER LETTER WORDS : 250 WORDS OR 1500 CHARACTER (whatever is minimum)

1. Start with:
   Kind Attention –
   Mr. {{executives}}
   M/s {{client_name}}

   Offer Ref: {{offer_ref}}; Date: {{letter_date}}

   Subject – Techno-Commercial Offer for {{project_title}}  

2. If there is only one executive, address them with:
   Dear {{first_exec_name}},
   If multiple executives, skip "Dear" and go directly to the content.
   Use Mr. for male and Ms. for female executives.

3. Opening paragraph (human way):
   - Acknowledge the invitation or requirement.
   - If invitation_date exists, mention it naturally.
   - If meeting_date exists, reference recent discussions or suggestions.
   - Wording must change between runs (not fixed sentences).

4. Body (human way):
   - Highlight Falcon’s analysis, solution evaluation, and technical proposal attachment.
   - Mention Falcon’s proven intralogistics technologies and experience.
   - Personalize with client_name.
   - Optionally mention project planning or timeline.

5. Closing (human way):
   - Reaffirm sender’s personal commitment.
   - Encourage the client to reach out for clarifications.
   - End with "Best Regards," followed by sender_name and sender_title.

Important:
- Do not exceed word/character limit.
- Keep tone formal, professional, and client-oriented.
- Do not copy exact sentences; rephrase wording across generations.
- The cover letter MUST sound human, natural and professional. It should be clear, authentic, and warm, without feeling robotic or overly formal.
- DO NOT ADD ANY EXTRA WORD OR INFO APART FROM THE COVER LETTER.
- Highlight the main system or project name in main body (not subject line) as bold style, use ** for Bold.
"""

USER_PROMPT_TEMPLATE = """
Use the following information to generate the cover letter:

client_name: {client_name}
project_title: {project_title}
offer_ref: {offer_ref}
letter_date: {letter_date}

executives (one per line, already with Mr./Ms. prefix):
{executives_block}

invitation_date: {invitation_date}
meeting_date: {meeting_date}

sender_name: {sender_name}
sender_title: {sender_title}

Return ONLY the cover letter text, without markdown code fences or extra commentary.
"""


def call_groq_cover_letter(
    client_name: str,
    project_title: str,
    offer_ref: str,
    letter_date_str: str,
    executives_block: str,
    invitation_date: str,
    meeting_date: str,
    sender_name: str,
    sender_title: str,
) -> str:
    user_prompt = USER_PROMPT_TEMPLATE.format(
        client_name=client_name,
        project_title=project_title,
        offer_ref=offer_ref,
        letter_date=letter_date_str,
        executives_block=executives_block.strip() or "Not provided",
        invitation_date=invitation_date.strip() or "Not provided",
        meeting_date=meeting_date.strip() or "Not provided",
        sender_name=sender_name,
        sender_title=sender_title,
    )

    completion = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.3,
        max_tokens=800,
    )

    text = completion.choices[0].message.content.strip()
    if text.startswith("```"):
        parts = text.split("```")
        if len(parts) >= 2:
            text = parts[1].lstrip("text").lstrip("markdown").lstrip()
    return text.strip()


# ========= DOCX helpers =========
HEADER_PREFIXES = (
    "Kind Attention",
    "Mr.",
    "Ms.",
    "M/s",
    "Offer Ref:",
    "Subject –",
    "Subject -",
)


def add_markdown_line(doc: Document, line: str):
    """Paragraph with **bold** segments."""
    p = doc.add_paragraph()
    parts = line.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        if i % 2 == 1:
            run.bold = True


def shade_cell(cell, color_hex: str = "D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_boxed_text(doc: Document, text: str, font_size: int = 16, bold: bool = True):
    """Grey shaded single-cell table with centered text (for 3D-ish header boxes)."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, "D9D9D9")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    # spacing before/after a bit smaller
    p.space_before = Pt(6)
    p.space_after = Pt(6)
    return table


def add_centered_upload_image(doc: Document, uploaded_file, width_in: float = 6.0):
    if not uploaded_file:
        return
    img_stream = BytesIO(uploaded_file.getvalue())
    p = doc.add_paragraph()
    r = p.add_run()
    r.add_picture(img_stream, width=Inches(width_in))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def write_cover_letter_to_doc(doc: Document, letter_text: str, sender_name: str, sender_title: str):
    lines = letter_text.splitlines()
    for line in lines:
        stripped = line.rstrip()

        if stripped == "":
            doc.add_paragraph("")
            continue

        if any(stripped.startswith(pref) for pref in HEADER_PREFIXES):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            continue

        lower = stripped.lower()
        if lower.startswith("best regards"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            continue

        if stripped.strip() in {sender_name.strip(), sender_title.strip()}:
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            continue

        add_markdown_line(doc, stripped)


def build_full_docx(
    letter_text: str,
    sender_name: str,
    sender_title: str,
    project_title: str,
    offer_ref: str,
    contact_name: str,
    contact_email: str,
    contact_phone: str,
    uploaded_image,
) -> BytesIO:
    doc = Document()

    # ---- Page 1: Cover letter ----
    write_cover_letter_to_doc(doc, letter_text, sender_name, sender_title)

    # ---- Page 2: Front page layout ----
    doc.add_page_break()

    # Top box: "Response to RFP for"
    add_boxed_text(doc, "Response to RFP for", font_size=18, bold=True)

    # Project name box
    if project_title:
        add_boxed_text(doc, project_title, font_size=18, bold=True)

    # Proposal reference box
    if offer_ref:
        add_boxed_text(doc, f"Proposal Reference – {offer_ref}", font_size=14, bold=True)

    # Some vertical spacing
    doc.add_paragraph("")

    # Layout image
    add_centered_upload_image(doc, uploaded_image, width_in=6.5)

    # More spacing
    doc.add_paragraph("")

    # Contact box at bottom
    contact_lines = [
        "Falcon Autotech Private Limited",
        "Plot No. 87, Sector Ecotech-1, Extention-1, Greater Noida, Uttar Pradesh 201308.",
        "",
        f"Contact – {contact_name}",
        "Assistant Manager",
        f"Mob - {contact_phone}",
        contact_email,
    ]
    contact_text = "\n".join(contact_lines)
    table = add_boxed_text(doc, contact_text, font_size=11, bold=False)
    # make contact text paragraphs centered
    cell = table.rows[0].cells[0]
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ========= Streamlit UI =========
st.set_page_config(page_title="Falcon Cover Letter + Front Page", page_icon="📄")
st.title("Techno-Commercial Cover Letter + Front Page Generator")

col1, col2 = st.columns(2)
with col1:
    client_name = st.text_input("Client name", value="Amazon")
    project_title = st.text_input("Project / system name", value="Loop CBS 48K + Linear CBS 5.4K PPH Throughput")
    offer_ref = st.text_input("Offer reference", value="F24-00524")

with col2:
    letter_date = st.date_input("Letter date", value=date.today())
    invitation_date = st.text_input("Invitation date (optional)", value="")
    meeting_date = st.text_input("Meeting / workshop date(s) (optional)", value="")

executives_text = st.text_area(
    "Executives (one per line, include Mr./Ms.)",
    value="Mr. Rahul Didwani\nMr. Vinayak Garg",
    height=80,
)

st.markdown("---")
st.subheader("Sender / Contact Details")

col3, col4 = st.columns(2)
with col3:
    sender_name = st.text_input("Sender name (for cover letter)", value="Sandeep Bansal")
    sender_title = st.text_input("Sender title", value="Chief Business Officer")
with col4:
    contact_name = st.text_input("Front-page contact name", value="Sanyog Pratap Singh")
    contact_phone = st.text_input("Front-page mobile", value="+91 8750052591")
    contact_email = st.text_input("Front-page email", value="Sanyog.Singh@falconautotech.com")

uploaded_image = st.file_uploader("Upload layout image for front page (PNG/JPG)", type=["png", "jpg", "jpeg"])

st.markdown("---")

generate = st.button("Generate DOCX (Cover Letter + Front Page)")

if generate:
    try:
        with st.spinner("Generating cover letter via Groq and building DOCX..."):
            letter_date_str = letter_date.strftime("%d-%m-%Y")
            letter_text = call_groq_cover_letter(
                client_name=client_name.strip(),
                project_title=project_title.strip(),
                offer_ref=offer_ref.strip(),
                letter_date_str=letter_date_str,
                executives_block=executives_text,
                invitation_date=invitation_date,
                meeting_date=meeting_date,
                sender_name=sender_name.strip(),
                sender_title=sender_title.strip(),
            )

            st.session_state["cover_letter_text"] = letter_text

            docx_buffer = build_full_docx(
                letter_text=letter_text,
                sender_name=sender_name.strip(),
                sender_title=sender_title.strip(),
                project_title=project_title.strip(),
                offer_ref=offer_ref.strip(),
                contact_name=contact_name.strip(),
                contact_email=contact_email.strip(),
                contact_phone=contact_phone.strip(),
                uploaded_image=uploaded_image,
            )

        st.success("DOCX generated.")

        st.subheader("Cover Letter (preview)")
        st.text_area("Generated cover letter text", value=st.session_state["cover_letter_text"], height=260)

        st.download_button(
            label="Download Cover Letter + Front Page (.docx)",
            data=docx_buffer,
            file_name=f"Proposal_{client_name or 'Client'}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except Exception as e:
        st.error(f"Error: {e}")






