import streamlit as st
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Exclusions Builder", page_icon="📄")

st.title("Exclusions Section Builder")

# ------------------------
# Fixed intro (always added)
# ------------------------
INTRO_TEXT = (
    "The scope of supply includes all parts which are defined in the Supplier’s quotation.\n"
    "All other parts which are not defined in the Supplier’s quotation do not belong to the Supplier’s "
    "scope of supply and are excluded. The following parts are also excluded:"
)

# ------------------------
# Fixed exclusions (always included)
# ------------------------
fixed_exclusions = [
    "Construction Power",
    "Building infrastructure; building structure, doors, fire exits, levelling devices, "
    "building extinguisher and fire alarm system, building heating and lighting system.",
    "Electrical power supply and wiring to the main control cabinets.",
    "UPS for Controls and Drives",
    "Network cabling up to the main server rack.",
    "Intermediate wiring to parts which are to be supplied by the Purchaser/others.",
    "Emergency/Uninterruptable power supply.",
    "Fire-alarm and fire protection devices.",
    "Traffic and route markings.",
    "Laydown area / unloading and laydown area.",
    "Ram protection devices.",
    "Cat walks, bridges, maintenance aisles and platforms.",
    "All kind of network incl. Local Area Network (LAN/WLAN), exceeding the scope described in Scope of Supply.",
    "Any kind of civil work.",
    "Any adjustment of the Supplier’s scope of supply to local rules and regulations.",
    "X-Ray machines.",
    "Roller cages / pallets.",
    "Simulation and 3D animation of the sorter system.",
    "Interface with other equipment not specified in this offer.",
    "Provision of facilities for the control room (furniture, air conditioning, heating, etc.).",
    "The supply and installation of fencing around the different corridors.",
    "Any item specifically indicated as not forming part of the subject matter of the Seller's supply in the offer documentation.",
]

# ------------------------
# Variable exclusions (user-selectable)
# ------------------------
variable_exclusions = [
    "Server PC / server system.",
    "SCADA / PC for SCADA.",
    "Workstations.",
    "Cabling from server room to Falcon control panel.",
    "Mobile carts.",
    "Collection trolleys / collection trolleys below chutes.",
    "Collection bins.",
    "Pallets at chutes.",
    "Pallets / hand-held terminals for secondary sorting.",
    "Steel works.",
    "Steel works – if not specified.",
    "Mezzanine & staircase.",
    "Mezzanine & staircase not mentioned in BOM.",
    "Maintenance platform / lift required for maintenance activity.",
    "Safety fencing / safety fencing not shown in layout.",
    "HPT/BOPT/Forklift/Hydra/Scaffoldings required for installation.",
    "Stress free mats.",
    "Insulation mats.",
    "Fans at chutes & inducts.",
    "Lighting around chutes / inducts.",
    "Irregular’s provision.",
    "UPS power (separate UPS supply).",
    "CE declaration of conformity.",
]



st.markdown("---")

st.subheader("Project-specific exclusions (tick to include)")
selected_variable = []
for idx, item in enumerate(variable_exclusions):
    if st.checkbox(item, key=f"var_{idx}"):
        selected_variable.append(item)

st.markdown("---")

def build_docx(fixed_items, variable_items):
    doc = Document()

    # Heading 2: Exclusions
    doc.add_heading("Exclusions", level=2)

    # Intro paragraph
    doc.add_paragraph(INTRO_TEXT)

    # Bullet list: fixed first, then variable
    for item in fixed_items + variable_items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

st.subheader("Generate Exclusions DOCX")

if st.button("Generate DOCX"):
    docx_buffer = build_docx(fixed_exclusions, selected_variable)
    st.download_button(
        label="Download Exclusions Section (.docx)",
        data=docx_buffer,
        file_name="Exclusions_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
