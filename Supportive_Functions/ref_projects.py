import os
from io import BytesIO

import streamlit as st
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="Reference Projects – Section Builder", page_icon="📌")

st.title("Reference Projects – DOCX Section Generator")
st.write(
    "Click **Generate** to create the fixed ‘Reference Projects’ section as a DOCX file, "
    "formatted similar to your proposal template (one project per page with site pictures)."
)

# ---------------------------------------------------
# CONFIG – update these paths to your actual images
# ---------------------------------------------------
IMG_REF_PROJ1 = r"FIXED_IMAGE\\proj1.PNG"
IMG_REF_PROJ2 = r"FIXED_IMAGE\\proj2.PNG"
IMG_REF_PROJ3 = r"FIXED_IMAGE\\proj3.PNG"
IMG_REF_PROJ4 = r"FIXED_IMAGE\\proj4.PNG"
IMG_REF_PROJ5 = r"FIXED_IMAGE\\proj5.PNG"

IMAGE_WIDTH = 6.0  # inches


def add_centered_image(doc: Document, img_path: str, width_in: float = IMAGE_WIDTH):
    """Add image centered if file exists."""
    if not img_path or not os.path.exists(img_path):
        return
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_reference_projects_docx() -> BytesIO:
    """Build full 'Reference Projects' section and return DOCX as BytesIO."""
    doc = Document()

    # ----------------- Top heading + intro bullets (Page 1) -----------------
    heading = doc.add_paragraph("Reference Projects")
    heading.runs[0].bold = True
    heading.style = "Heading 2"

    intro = (
        "Falcon has a strong legacy in Warehousing Automation solutions and references-"
    )
    doc.add_paragraph(intro)

    bullets_intro = [
        "Expertise in Shipment Sortation, Piece Picking and Handling, Case Picking and Handling.",
        "Lifecycle services (maintenance, spares supply chain, support).",
        "Full in-house expertise (Hardware/Software).",
        "Turn-key tailored solutions.",
        "The references list presented below focuses on Sortation Solution –",
    ]
    for text in bullets_intro:
        doc.add_paragraph(text, style="List Bullet")

    # ----------------- Project 1 (still on Page 1) -----------------
    p = doc.add_paragraph("5.1 Project 1- (CEP Client, India)")
    p.style = "Heading 3"

    doc.add_paragraph(
        "The system is equipped with two fully automated and interconnected sub-systems. "
        "Sub-System 1 is designed for handling large B2B boxes and E-commerce shipment bags while "
        "Sub-System 2 is designed to handle small E-commerce packages."
    )

    doc.add_paragraph("Solution Specifications –")
    spec1 = [
        "48,000 PPH (Double Deck CBS – Shipment Sorter).",
        "17,000 PPH (Double Deck CBS – Bag Sorter).",
        "Building Size: 700,000 Sq. Ft.",
    ]
    for t in spec1:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Key Technology Modules –")
    ktm1 = [
        "2 Sets of Double Decker CBS Sorters.",
        "Mezzanine Structures.",
        "Automated Singulators.",
        "Fully Automatic Inductions.",
        "Semi-Automatic Inductions.",
        "Telescopic Belt Conveyors.",
        "PVC Belt Conveyors.",
        "Modular Belt Conveyors.",
        "Spiral Chutes with Braking Rollers.",
        "5-Sided Scanning Tunnels.",
        "High Speed Weighing Conveyors.",
        "Direct Bagging Chutes.",
        "Put to Light Chutes.",
        "Volume Distribution Systems.",
        "High Availability Server Systems.",
        "WCS.",
    ]
    for t in ktm1:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Site Pictures –")
    add_centered_image(doc, IMG_REF_PROJ1)

    doc.add_page_break()

    # ----------------- Project 2 (Page 2) -----------------
    p = doc.add_paragraph("5.2 Project 2- (Client – E-Commerce, India)")
    p.style = "Heading 3"

    doc.add_paragraph("Use Case – Destination sorting of packed shipments.")

    doc.add_paragraph(
        "In 2019, the client was looking for a potential automation partner for design and development of a "
        "new automated sortation system for B2C shipments. The system needed to provide maximum uptime with "
        "reduced dependency on skilled manpower and better space optimization. "
        "The customer chose Falcon Autotech based on its unique design that addressed these pain points, "
        "its capability for seamless WMS integration, and its life cycle support services."
    )

    doc.add_paragraph("Solution Specifications –")
    spec2 = [
        "Throughput: 27,600 PPH.",
        "End Destinations: 410 Direct Outputs.",
        "Building Size: 200,000 Sq. Ft.",
    ]
    for t in spec2:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Key Technology Modules –")
    ktm2 = [
        "Bulk Infeed Conveyors.",
        "ARB based Volume Distribution System.",
        "Integrated Presort System.",
        "Irregular Ejection System.",
        "Automatic Induct Lines.",
        "Automatic Barcode Scanner with Image Capture.",
        "Automatic Weight & Volume Measurement System.",
        "Linear Cross Belt Sorter.",
        "Smart Sliding Chutes for Direct Bagging and Cage Sorting.",
        "Bag Take-out System.",
        "WCS Software System.",
    ]
    for t in ktm2:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Site Pictures –")
    add_centered_image(doc, IMG_REF_PROJ2)

    doc.add_page_break()

    # ----------------- Project 3 (Page 3) -----------------
    p = doc.add_paragraph("5.3 Project 3- (Client – E-Commerce, India)")
    p.style = "Heading 3"

    doc.add_paragraph("Use Case – Destination sorting of packed shipments.")

    doc.add_paragraph(
        "The customer chose Falcon Autotech based on its unique design, its ability to integrate seamlessly "
        "with the WMS, and its strong life cycle support services."
    )

    doc.add_paragraph("Solution Specifications –")
    spec3 = [
        "Throughput: 24,000 PPH.",
        "End Destinations: 40 Collection Type Chutes.",
    ]
    for t in spec3:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Key Technology Modules –")
    ktm3 = [
        "Bulk Infeed Conveyors.",
        "ARB based Volume Distribution System.",
        "Irregular Ejection System.",
        "Automatic Induct Lines.",
        "Automatic Barcode Scanner with Image Capture.",
        "Automatic Weight & Volume Measurement System.",
        "Linear Cross Belt Sorter.",
        "Smart Collection Type Chutes.",
        "Bag Take-out System.",
        "WCS Software System.",
    ]
    for t in ktm3:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Site Pictures –")
    add_centered_image(doc, IMG_REF_PROJ3)

    doc.add_page_break()

    # ----------------- Project 4 (Page 4) -----------------
    p = doc.add_paragraph("5.4 Project 4- (CEP Client, UK)")
    p.style = "Heading 3"

    doc.add_paragraph(
        "This solution is designed to handle a volume of 7,200 shipments per hour. "
        "The system is equipped with three infeed conveyors integrated with an automatic label applicator "
        "before shipments enter the sortation system. Shipments are sorted using Falcon’s Loop Cross Belt Sorter "
        "equipped with automatic barcode scanning, dimensioning, weighing, and image capture capabilities. "
        "The sorter is installed on the mezzanine floor and sorts directly to 58 end destinations."
    )

    doc.add_paragraph("Solution Specifications –")
    spec4 = [
        "Throughput: 7,200 PPH.",
        "End Destinations: 58 Nos.",
    ]
    for t in spec4:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Key Technology Modules –")
    ktm4 = [
        "Powered Belt Conveyors.",
        "Automatic Induct Lines.",
        "Automatic Barcode Scanner with Image Capture.",
        "Automatic Weight & Volume Measurement System.",
        "Loop Cross Belt Sorter.",
        "WCS Software System.",
    ]
    for t in ktm4:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Site Picture –")
    add_centered_image(doc, IMG_REF_PROJ4)

    doc.add_page_break()

    # ----------------- Project 5 (Page 5) -----------------
    p = doc.add_paragraph("5.5 Project 5- (CEP Client, Sydney)")
    p.style = "Heading 3"

    doc.add_paragraph(
        "This solution is designed for handling a throughput of 16,000 shipments per hour with the help of "
        "Falcon’s Loop Cross Belt Sorter. The system consists of two feeding zones with a total of ten feedlines. "
        "Sorter design enables van drivers to directly drop shipments at the dock doors. It has a total of 369 end "
        "destinations achieved through a combination of direct drops and PTLs. The system is integrated with "
        "five-side automatic barcode scanning, weight and volume measurement, and automatic detection of "
        "oversize and overweight shipments."
    )

    doc.add_paragraph("Solution Specifications –")
    spec5 = [
        "Throughput: 16,000 PPH.",
        "End Destinations: 369 Nos.",
    ]
    for t in spec5:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Key Technology Modules –")
    ktm5 = [
        "Powered Belt Conveyors.",
        "2 Induct Zones.",
        "5-side Automatic Barcode Scanner.",
        "Automatic Weight & Volume Measurement System.",
        "Automatic Detection of Oversize Shipments.",
        "Loop Cross Belt Sorter.",
        "WCS Software System.",
    ]
    for t in ktm5:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_paragraph("Site Picture –")
    add_centered_image(doc, IMG_REF_PROJ5)

    # -------------- Save to buffer --------------
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


docx_buffer = None
if st.button("Generate Reference Projects DOCX"):
    with st.spinner("Building Reference Projects section..."):
        docx_buffer = build_reference_projects_docx()

if docx_buffer:
    st.download_button(
        label="Download Reference Projects Section (.docx)",
        data=docx_buffer,
        file_name="Reference_Projects_Section.docx",
        mime=(
            "application/"
            "vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
