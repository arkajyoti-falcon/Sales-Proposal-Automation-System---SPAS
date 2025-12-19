import re
from io import BytesIO

import streamlit as st
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.shared import Pt


# --------------------------------------------------------------------
# 1. Master glossary list (Term -> Description)
# --------------------------------------------------------------------
GLOSSARY_ENTRIES = [
    ("RFQ", "Request For Quotation"),
    ("RFP", "Request For Proposal"),
    ("PPH", "Parcels / Shipments Per Hour"),
    ("ARB", "Actuated Roller Balls"),
    ("DBO", "Damaged Barcode"),
    ("VDS", "Volume Distribution System"),
    ("ICR", "Intelligent Character Recognition"),
    ("MEZZ", "Mezzanine"),
    ("LIM", "Linear Induction Motor"),
    ("LSM", "Linear Synchronous Motor"),
    ("FWD", "Friction Wheel Drive"),
    ("ECDS", "Empty Carrier Detection System"),
    ("AC", "Alternating Current"),
    ("DC", "Direct Current"),
    ("PLC", "Programmable Logic Controller"),
    ("IT", "Information Technology"),
    ("BOQ", "Bill Of Quantity"),
    ("I/O", "Input/ Output"),
    ("PDP", "Power Distribution Panel"),
    ("PC", "Personal Computer"),
    ("UPS", "Uninterrupted Power Supply"),
    ("CBS", "Cross Belt Sorter"),
    ("MDR", "Motor Driven Roller"),
    ("IPP", "Individual Productivity Potential"),
    ("VM", "Virtual Machine"),
    ("MENA", "Middle East North Africa"),
    ("FOC", "Free of Cost"),
    ("CEP", "Courier Express Parcel"),
    ("DAP", "Design Approval Phase"),
    ("DNF", "Data Not Found"),
    ("LGM", "Logic Mismatch"),
    ("RTVC", "Real Time Video Coding"),
    ("NL Shipments", "Non-Large Shipments"),
    ("SL Shipments", "Semi-Large Shipments"),
    ("NO(S)", "Piece(s)"),
]


# --------------------------------------------------------------------
# 2. Utilities
# --------------------------------------------------------------------
def extract_full_text_from_docx(doc: Document) -> str:
    """Concatenate all paragraph and table text from the DOCX."""
    parts = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts)


def find_terms_in_text(text: str):
    """
    Return list of (term, description) that actually appear in the text.
    Matching is done as a whole word, case-sensitive, to avoid accidental
    matches of 'it' vs 'IT'.
    """
    found = []
    for term, desc in GLOSSARY_ENTRIES:
        # Build safe regex – whole word / token
        pattern = r"(?<![A-Za-z0-9])" + re.escape(term) + r"(?![A-Za-z0-9])"
        if re.search(pattern, text):
            found.append((term, desc))

    # keep original order, no duplicates
    seen = set()
    unique = []
    for t, d in found:
        if t not in seen:
            seen.add(t)
            unique.append((t, d))
    return unique


def build_glossary_doc(terms):
    """
    Create a new DOCX with '1. Glossary' heading and a table:
    S. No. | Term | Description
    """
    doc = Document()

    # Heading
    heading = doc.add_heading("1. Glossary", level=1)
    heading.alignment = 0  # left

    # Table
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    headers = ["S. No.", "Term", "Description"]
    for i, text in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        p.alignment = 1  # center
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Rows
    for idx, (term, desc) in enumerate(terms, start=1):
        row_cells = table.add_row().cells

        # S. No.
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(str(idx))
        r0.bold = True
        p0.alignment = 1  # center
        row_cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Term
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(term)
        r1.bold = True
        p1.alignment = 1
        row_cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Description
        p2 = row_cells[2].paragraphs[0]
        p2.add_run(desc)
        row_cells[2].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Optional: unify font size a bit
        for cell in row_cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)

    return doc


# --------------------------------------------------------------------
# 3. Streamlit app
# --------------------------------------------------------------------
st.title("Glossary Extractor for Falcon Proposals")

st.write(
    "Upload the **final proposal DOCX** (without glossary). "
    "The app will scan the content for known abbreviations and "
    "generate a separate **Glossary.docx** with a table like your template."
)

uploaded = st.file_uploader("Upload proposal (.docx)", type=["docx"])

if uploaded is not None:
    # Read proposal DOCX
    proposal_doc = Document(uploaded)
    full_text = extract_full_text_from_docx(proposal_doc)

    # Detect which terms actually appear
    detected_terms = find_terms_in_text(full_text)

    if not detected_terms:
        st.warning("No glossary terms found in this document based on the current list.")
    else:
        st.success(f"Found {len(detected_terms)} glossary term(s) in the proposal.")

        # Show a preview table in Streamlit
        st.subheader("Detected Glossary Terms (preview)")
        st.table(
            {
                "S. No.": list(range(1, len(detected_terms) + 1)),
                "Term": [t for t, _ in detected_terms],
                "Description": [d for _, d in detected_terms],
            }
        )

        # Build DOCX and offer download
        glossary_doc = build_glossary_doc(detected_terms)
        bio = BytesIO()
        glossary_doc.save(bio)
        bio.seek(0)

        st.download_button(
            label="Download Glossary DOCX",
            data=bio,
            file_name="Glossary.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
