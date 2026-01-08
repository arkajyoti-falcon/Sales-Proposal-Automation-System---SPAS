import streamlit as st
import pandas as pd
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Key Components Make", page_icon="🧩")

st.title("Key Components Make – Section Builder")
st.write(
    "This section creates a single consolidated table of key components and their proposed makes, "
    "based on your standard proposals. You can adjust makers and then export to DOCX."
)

# --------- Default consolidated mapping (derived from your examples) ---------
DEFAULT_ROWS = [
    {"Items": "Belts", "Make": "Forbo / Derco / Habasit"},
    {"Items": "Rollers", "Make": "Falcon"},
    {"Items": "Cross Belt Carriers", "Make": "Falcon"},
    {"Items": "Linear Motors (LIM / LSM / Linear Induction)", "Make": "Falcon / SEW / FWD (as applicable)"},
    {"Items": "Feed Line Motors", "Make": "Falcon"},
    {"Items": "Volume / Barcode Scanners", "Make": "SICK / Cognex / Similar"},
    {"Items": "Weighing Scales", "Make": "Bizerba / Mettler Toledo / Equivalent"},
    {"Items": "Encoders", "Make": "SICK / Falcon"},
    {"Items": "Sensors", "Make": "SICK / Leuze / P&F"},
    {"Items": "PLC", "Make": "Siemens / Omron"},
    {"Items": "Control Panels", "Make": "Rittal / BCH"},
    {"Items": "VFDs", "Make": "Siemens / Lenze / AB / Omron"},
    {"Items": "Cables", "Make": "LAPP / Equivalent"},
    {"Items": "Switch Gear", "Make": "Schneider / Equivalent"},
    {"Items": "Bearings", "Make": "NTN / SKF / Equivalent"},
    {"Items": "Power Transmission Systems", "Make": "Vahle"},
    {"Items": "HMIs", "Make": "Siemens / Omron"},
    {"Items": "MDR", "Make": "Pulse / Itoh Denki"},
    {"Items": "Data Transmission System", "Make": "Siemens"},
]

# Initialise session state once
if "key_components_df" not in st.session_state:
    st.session_state["key_components_df"] = pd.DataFrame(DEFAULT_ROWS)

st.subheader("Edit Key Components Make Table")
edited_df = st.data_editor(
    st.session_state["key_components_df"],
    num_rows="dynamic",
    use_container_width=True,
    key="key_components_editor",
)
st.session_state["key_components_df"] = edited_df


def build_key_components_docx(df: pd.DataFrame) -> BytesIO:
    doc = Document()

    # Heading 2
    doc.add_heading("Key Components Make", level=2)

    # Table: Items | Make
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    hdr[0].text = "Items"
    hdr[1].text = "Make"

    for _, row in df.iterrows():
        r = table.add_row().cells
        r[0].text = str(row.get("Items", ""))
        r[1].text = str(row.get("Make", ""))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


st.markdown("---")
st.subheader("Download DOCX")

docx_buf = build_key_components_docx(st.session_state["key_components_df"])
st.download_button(
    label="Generate & Download Key Components Make (.docx)",
    data=docx_buf,
    file_name="Key_Components_Make.docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
)
