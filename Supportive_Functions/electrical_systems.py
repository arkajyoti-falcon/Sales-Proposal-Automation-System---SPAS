import os
from io import BytesIO

import streamlit as st
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="Electrical System – Section Builder", page_icon="⚡")

st.title("Electrical System – DOCX Section Generator")

st.write(
    "Click the button below to generate a ready-to-insert DOCX section "
    "for the *Electrical System* with fixed content and reference images."
)

# -------------------------------------------------------------------
# CONFIG – image paths (update to match your actual folder structure)
# -------------------------------------------------------------------
IMG_PDP = "FIXED_IMAGE/elec1.PNG"              # 15.1 – Reference Picture of Power Distribution Panel
IMG_MAIN_CTRL = "FIXED_IMAGE/elec2.PNG"  # 15.2 – Main Control Panel (Reference)
IMG_INDUCT_CTRL = "FIXED_IMAGE/elec3.PNG"  # 15.3 – Induct Stations Control Panel (Reference)

IMAGE_WIDTH_INCHES = 5.5  # reasonable width for A4 page


def add_centered_image(doc: Document, path: str, width_in: float = IMAGE_WIDTH_INCHES):
    """
    Safely add a centered image if the file exists.
    If the file does not exist, it silently skips it.
    """
    if not path or not os.path.exists(path):
        return

    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_electrical_system_docx() -> BytesIO:
    """Build the 'Electrical System' section as a DOCX and return as BytesIO."""
    doc = Document()

    # Main section heading
    doc.add_heading("Electrical System", level=2)

    # Intro paragraph (from PDF)
    doc.add_paragraph(
        "Main power supply will supply Falcon’s PDP (Power Distribution Panels) electrical cabinets. "
        "PDP cabinets supply the entire system via secondary cabinets:"
    )

    # Bullet list of secondary cabinets
    doc.add_paragraph("Main Control Cabinet", style="List Bullet")
    doc.add_paragraph("Induct Control Panels", style="List Bullet")
    doc.add_paragraph("Remote Cabinets for Sorter I/O", style="List Bullet")
    doc.add_paragraph("Scanner Control cabinets", style="List Bullet")

    # 15.1 – Reference Picture of Power Distribution Panel
    doc.add_heading("Reference Picture of Power Distribution Panel", level=3)
    add_centered_image(doc, IMG_PDP)

    # 15.2 – Main Control Panel (Reference)
    doc.add_heading("Main Control Panel (Reference)", level=3)
    add_centered_image(doc, IMG_MAIN_CTRL)

    # 15.3 – Induct Stations Control Panel (Reference)
    doc.add_heading("Induct Stations Control Panel (Reference)", level=3)
    add_centered_image(doc, IMG_INDUCT_CTRL)

    # Engines
    p = doc.add_paragraph()
    p.add_run("Engines").bold = True

    doc.add_paragraph(
        "Three-phase alternating current motors (Induction) will be used through a frequency converter. "
        "The engines will be coupled with a converter to improve consumption and reduce the carbon footprint."
    )
    doc.add_paragraph(
        "All motors will have appropriate IP ratings."
    )

    # Sensors
    p = doc.add_paragraph()
    p.add_run("Sensors").bold = True

    doc.add_paragraph(
        "The sensors will be supplied, standardized by type, with connector, with a cable length suitable for "
        "easy extraction, suitably protected from possible impacts."
    )

    # Control command
    p = doc.add_paragraph()
    p.add_run("Control command").bold = True

    doc.add_paragraph(
        "The proposed solution is based on SIEMENS Programmable Logic Controller technology (PLC) platform. "
        "The entire system will be logically divided into Zones (Sorter / Feed Line / Loop), each managed by a PLC. "
        "The planned primary communication protocol is going to be ProfiNet."
    )

    # Conveyor interface
    p = doc.add_paragraph()
    p.add_run("Conveyor interface").bold = True

    doc.add_paragraph(
        "The frequency converter of each conveyor allows the acquisition of the signals of the "
        "sensors/actuators/GIOs associated with it (e.g. conveyor end detection photocells, blockage detection "
        "photocells). Each frequency converter will be connected in series by means of the ProfiNet field bus."
    )

    # Save to buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# -----------------------------
# Streamlit UI
# -----------------------------
docx_buffer = None

if st.button("Generate Electrical System DOCX"):
    with st.spinner("Building Electrical System section..."):
        docx_buffer = build_electrical_system_docx()

if docx_buffer:
    st.download_button(
        label="Download Electrical System Section (.docx)",
        data=docx_buffer,
        file_name="Electrical_System_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
