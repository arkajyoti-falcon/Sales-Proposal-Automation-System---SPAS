import os
from io import BytesIO

import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ----------------- CONFIG -----------------
STATIC_ABOUT_DIR = r"Static_AboutCompany"   # base folder for all 1.png ... 12.png


# ----------------- HELPERS -----------------
def add_section_heading(doc: Document, text: str) -> None:
    """
    Adds a section heading similar to your proposal:
    bold + underlined, left-aligned.
    """
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.underline = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_body_text(doc: Document, text: str) -> None:
    """
    Adds a normal body paragraph with Calibri 11 and left alignment.
    """
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_center_image(doc: Document, filename: str, width_in: float = 6.0) -> None:
    """
    Adds an image centered on the page (if it exists).
    """
    path = os.path.join(STATIC_ABOUT_DIR, filename)
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ----------------- MAIN BUILDER -----------------
def build_company_profile(doc: Document) -> None:
    """
    Populate the given `doc` with the Company Profile content.
    Mirrors structure/text from your source + static PNGs.
    """

    # ---------- Page 1 ----------
    add_section_heading(doc, "Company Profile")

    top_text = (
        "Falcon Autotech (Falcon) is a global intralogistics automation solutions company. "
        "With over 10 years of experience, Falcon has worked with some of the most innovative "
        "brands in E-Commerce, CEP, Fashion, Food/FMCG, Auto and Pharmaceutical Industries. "
        "With our proprietary software and robust hardware integration capabilities, Falcon designs, "
        "manufactures, supplies, implements, and maintains world-class warehouse automation systems globally. "
        "Falcon’s strong research and development team and the continuous focus on innovation reflect our strong "
        "solution line around Sortation, Robotics, Conveying, Vision Systems and IOT. "
        "Falcon has done over 1,800 installations across 15 countries on four continents."
    )
    add_body_text(doc, top_text)

    add_center_image(doc, "1.png", width_in=6)

    bottom_text = (
        "Falcon Autotech is currently among the top 15 intralogistics automation companies; "
        "our vision is to become a top 10 intralogistics automation company in our focused product lines."
    )
    add_body_text(doc, bottom_text)

    add_center_image(doc, "2.png", width_in=5)

    doc.add_page_break()

    # ---------- Page 2 ----------
    top_text2 = (
        "The team started out in 2004 solving special purpose automation problems for clients and later "
        "established Falcon Autotech in 2012 with a strong focus on building a standard technology stack spanning "
        "across hardware, firmware, and software to tackle larger supply chain problems around warehouse "
        "automation and material handling. "
        "Over the decade, Falcon has made rapid strides and has carved out a niche in some of the world's most "
        "cutting-edge technologies: Sortation, Robotics, Conveying, Vision Systems and IOT."
    )
    add_body_text(doc, top_text2)

    add_center_image(doc, "3.png", width_in=6)

    bottom_text2 = (
        "As a leading player in the intralogistics automation space, Falcon continuously strives to improve the "
        "operational efficiencies and accuracies for its clients through its domain knowledge and experience, in "
        "addition to its wide range of products and solutions. In order to live up to the high expectations set "
        "forth by our clients, the team at Falcon realizes the importance of taking up selective applications in "
        "focused industries and delivering world-class projects in return."
    )
    add_body_text(doc, bottom_text2)

    add_center_image(doc, "4.png", width_in=6)

    # ---------- Page 3 ----------
    add_center_image(doc, "5.png", width_in=6)

    bottom_text3 = (
        "Falcon Autotech has successfully delivered warehouse automation solutions based on smart and innovative "
        "combinations of the above product lines for effective materials handling, sortation and movement. "
        "The process is controlled in real-time by our in-house WCS applications. These solutions considerably "
        "reduce the need for manual operations, improve working conditions and ensure the highest accuracy of the "
        "entire process up to final delivery to the recipient.\n\n"
        "Over the last 10 years, Falcon has worked with some of the most innovative brands worldwide and has "
        "established long-standing partnerships. These brands are testimony to our strong focus on delivering "
        "superior customer satisfaction and offering end-to-end intralogistics solutions."
    )
    add_body_text(doc, bottom_text3)

    add_center_image(doc, "6.png", width_in=6)

    doc.add_page_break()

    # ---------- Page 4 ----------
    bottom_text4 = (
        "With over 1,800 installations, Falcon’s systems are used all over the globe. Falcon has a highly "
        "motivated team of 600+ employees supported by over 15 global partners who help us design, manufacture, "
        "deliver and maintain automation solutions worldwide."
    )
    add_body_text(doc, bottom_text4)

    add_center_image(doc, "7.png", width_in=6)

    add_section_heading(doc, "Customer Engagement Model")
    add_center_image(doc, "8.png", width_in=7)

    doc.add_page_break()

    # ---------- Page 5 ----------
    add_section_heading(doc, "Falcon’s Experience and Achievements in Sortation Space Globally")

    bullet_points = [
        "Ranked among Top 10 Sortation System Suppliers globally.",
        "Currently possess one of the world’s largest portfolios in sortation technologies (7 in-house technologies).",
        "Total installed capacity of 10 million shipments per day worldwide.",
        "Only company to be able to offer a fully integrated AMS.",
    ]

    for point in bullet_points:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(point)
        r.font.name = "Calibri"
        r.font.size = Pt(10)

    add_center_image(doc, "9.png", width_in=6)
    add_center_image(doc, "10.png", width_in=6)
    add_center_image(doc, "11.png", width_in=6)

    # p = doc.add_paragraph()
    # p.add_run().add_break()  # small spacing

    add_center_image(doc, "12.png", width_in=6)


def generate_company_profile_docx() -> BytesIO:
    """Create a fresh DOCX with the Company Profile section and return as BytesIO."""
    doc = Document()
    build_company_profile(doc)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ----------------- STREAMLIT UI -----------------
st.set_page_config(page_title="Company Profile – DOCX Generator", page_icon="🏢")

st.title("Company Profile – DOCX Section Generator")
st.write(
    "Click the button below to generate the full multi-page **Company Profile** section "
    "with static images and text, ready to insert into your proposal."
)

if st.button("Generate Company Profile DOCX"):
    with st.spinner("Building Company Profile document..."):
        docx_bytes = generate_company_profile_docx()

    st.download_button(
        label="Download Company Profile (.docx)",
        data=docx_bytes,
        file_name="Falcon_Company_Profile.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
