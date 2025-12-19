import streamlit as st
import pandas as pd
import json
import re
from io import BytesIO
from docx import Document
from groq import Groq
from dotenv import load_dotenv
import os

load_dotenv()

st.set_page_config(page_title="Price Sheet via Groq", page_icon="💰")

st.title("Commercial – Price Sheet Builder (Groq-powered)")
st.write(
    "Upload a costing Excel file. The app sends the 'Overall Costing' sheet to Groq and "
    "asks it to return a high-level Price Sheet summary like your proposal examples."
)

# -----------------------------
# CONFIG – Groq client & prompts
# -----------------------------
client = os.getenv("GROQ_API_KEY")

SYSTEM_PROMPT = """
You are a senior commercial analyst for warehouse automation projects.

You receive the contents of an Excel sheet called "Overall Costing" as raw CSV text.
This sheet may contain many detailed costing lines, intermediate totals, taxes, and notes.

Your job is to infer the HIGH-LEVEL "Price Sheet" summary used in proposals.

The high-level Price Sheet is a short table of a few summary lines (typically 3–15),
each corresponding to a major package/component of the solution with a single rolled-up price.
Do NOT list detailed items like small sub-components or line-by-line BOM;
only show the SUMMARY building blocks that a customer would see in the commercial section.

------------------------------------------------
EXAMPLES OF TARGET PRICE SHEETS (FOR REFERENCE)
------------------------------------------------

Example 1 –

Price List- Summary
S.NO   Package                    Price
1      Conveyors Package          ₹ 28,34,27,926
2      Cross Belt Sorter Package  ₹ 29,16,52,094
3      Destinations Package       ₹ 3,83,87,721
4      Services Package           ₹ 3,46,22,786
       Total                      ₹ 64,80,90,527

"Business Cooperation Agreement"
Discount for Delhivery  4.5%
Final Total             ₹ 61,89,26,453


Example 2 – 

Price Sheet
S. No   Component                               Price (USD)
1       Loop CBS + Inducts                      (included or price)
2       Infeed + Bagging Conveyors             $ 398,105
3       Output Chutes                          $ 219,944
4       Software Package & Integration         $ 26,302
5       Packaging & forwarding                 $ 3,523
6       Project Management + Supervision cost  $ 32,268
Total (USD)                                    $ 726,386

Followed by bullets:
• Inco-Terms- Ex-works, India (Greater Noida)
• Taxes: Extra as Applicable
• Price is valid for 30 days from the date of proposal.
(These bullets are not part of the high-level table, but the items+total are.)


---------------------------------------
TASK – WHAT YOU MUST RETURN
---------------------------------------

Use the raw 'Overall Costing' CSV to reconstruct ONLY the high-level summary, in the spirit of the examples above.

1) Identify the main commercial building blocks, such as:
   - Conveyors Package
   - Cross Belt Sorter Package
   - Destinations Package
   - Services Package
   - Loop CBS + Inducts
   - Infeed + Bagging Conveyors
   - Output Chutes
   - Software package / Software Packages & SCADA
   - Steelworks / Steelwork
   - PTL
   - Installation & Commissioning
   - Project Management and Engineering Charges
   - Packaging & forwarding / Packaging & Documentation
   - Freight, Warranty, Hotline, AMC packages
   or similar high-level components used to summarize the cost.

2) For each such high-level component, return:
   - s_no: integer starting from 1 in sequence
   - label: the package/component name in clean human-readable form
   - price: the final total price for that component AS A STRING,
            including currency symbol and formatting exactly as in the sheet
            (e.g. "₹ 28,34,27,926", "SAR 10,959,208", "$ 398,105", "€ 2,871,416", "Included in CBS Price").

   IMPORTANT:
   - Do NOT invent prices.
   - Use values that actually appear in the sheet.
   - If multiple detailed rows roll up into one package, use the rolled-up total that clearly corresponds to that package.
   - Prefer the same format as used for the final summary in the data, if visible.

3) If there is a grand "Total" (for the whole solution), also return:
   - total_row: { "label": "...", "price": "..." }
   For example: { "label": "Total", "price": "₹ 64,80,90,527" }.
   If no obvious total exists, set total_row to null.

4) If there are explicit discount and final total lines (like in the Delhivery example):
   - cooperation_label: e.g. "Business Cooperation Agreement"
   - discount_label: e.g. "Discount for Delhivery"
   - discount_value: e.g. "4.5%"
   - final_total_label: e.g. "Final Total"
   - final_total_value: e.g. "₹ 61,89,26,453"
   If not present, return them as null.

5) Also return:
   - currency: "INR", "SAR", "USD", "EUR", or "MIXED" if multiple currencies appear.
   - price_sheet_title: a short label like "19.1 Price List – Summary" or "20.1 Price Sheet"
                        if visible; otherwise null.

6) OUTPUT FORMAT (VERY IMPORTANT):

Return ONLY a single valid JSON object with this exact shape:

{
  "currency": "INR" | "SAR" | "USD" | "EUR" | "MIXED" | null,
  "price_sheet_title": "string or null",
  "items": [
    {
      "s_no": 1,
      "label": "Conveyors Package",
      "price": "₹ 28,34,27,926"
    },
    ...
  ],
  "total_row": {
    "label": "Total",
    "price": "₹ 64,80,90,527"
  } or null,
  "cooperation_label": "string or null",
  "discount_label": "string or null",
  "discount_value": "string or null",
  "final_total_label": "string or null",
  "final_total_value": "string or null"
}

Do NOT wrap the JSON or ```json``` in markdown.
Do NOT add explanations or commentary.
Just return the JSON object.
"""

USER_PROMPT_TEMPLATE = """
Below is the raw CSV export of the 'Overall Costing' sheet of an internal costing file.

Use it to construct the high-level Price Sheet summary as described in the instructions.

Raw CSV:
--------------------
{sheet_csv}
--------------------
"""


def call_groq_for_price_sheet(sheet_csv: str) -> dict:
    user_prompt = USER_PROMPT_TEMPLATE.format(sheet_csv=sheet_csv)

    completion = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.0,
    )

    raw = completion.choices[0].message.content.strip()

    # 1) If wrapped in ```json ... ``` or ``` ... ```, strip fences
    if raw.startswith("```"):
        parts = raw.split("```")
        if len(parts) >= 2:
            raw = parts[1]
            raw = raw.lstrip("json").lstrip()

    # 2) Extract from first '{' to last '}'
    start = raw.find("{")
    end = raw.rfind("}")
    if start == -1 or end == -1 or end < start:
        raise ValueError(f"Groq response does not contain a JSON object:\n{raw}")

    content = raw[start:end + 1]

    try:
        data = json.loads(content)
    except json.JSONDecodeError as e:
        raise ValueError(f"Groq response was not valid JSON: {e}\nExtracted content:\n{content}")

    return data


# ---------- helpers for discount formatting ----------

def parse_price_string(price_str: str):
    """
    Extract currency prefix and numeric value from a price string like:
    '₹ 647,398,701.71' or 'SAR 10,959,208' or '$ 398,105'
    Returns (prefix, value, decimals_count) or (None, None, 0) if parse fails.
    """
    if not price_str:
        return None, None, 0

    # Find first digit (or minus sign) in the string
    m = re.search(r"[-]?\d", price_str)
    if not m:
        return price_str.strip(), None, 0

    prefix = price_str[:m.start()].strip()
    numeric_part = price_str[m.start():].strip()

    # Strip non-digit separators except '.' for decimals
    digits_only = "".join(ch for ch in numeric_part if ch.isdigit() or ch == ".")
    if digits_only == "":
        return prefix, None, 0

    # Count decimals
    decimals_count = 0
    if "." in digits_only:
        decimals_count = len(digits_only.split(".")[1])

    try:
        value = float(digits_only)
    except ValueError:
        return prefix, None, decimals_count

    return prefix, value, decimals_count


def format_indian_number(value: float, decimals: int) -> str:
    """
    Format a number with Indian-style digit grouping.
    e.g. 647398701.71 -> "64,73,98,701.71" when decimals=2
    """
    if decimals > 0:
        s = f"{value:.{decimals}f}"
    else:
        s = f"{int(round(value))}"

    if "." in s:
        int_part, frac = s.split(".")
    else:
        int_part, frac = s, None

    # Indian grouping
    if len(int_part) > 3:
        last3 = int_part[-3:]
        head = int_part[:-3]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        int_formatted = ",".join(groups + [last3])
    else:
        int_formatted = int_part

    if frac and decimals > 0:
        return int_formatted + "." + frac
    else:
        return int_formatted


def apply_bca_discount_to_price_data(price_data: dict, discount_percent: float) -> str | None:
    """
    Apply BCA discount on total_row.price and return discounted price string.
    Does not modify original total_row; just returns the final_total string.
    """
    total_row = price_data.get("total_row")
    if not total_row:
        return None

    price_str = total_row.get("price")
    prefix, value, decimals = parse_price_string(price_str)
    if value is None:
        return None

    discounted = value * (1 - discount_percent / 100.0)
    formatted_number = format_indian_number(discounted, decimals)
    if prefix:
        return f"{prefix} {formatted_number}"
    else:
        return formatted_number


# ---------- DOCX builder ----------

def build_price_sheet_docx(
    price_data: dict,
    apply_bca: bool,
    payment_terms: list[dict],
    discount_percent: float = 4.5,
) -> BytesIO:
    """Create a DOCX file from the structured price sheet data + payment terms."""
    doc = Document()

    # Heading 2 – Price Sheet
    doc.add_heading("Price Sheet", level=2)

    title = price_data.get("price_sheet_title") or ""
    if title:
        doc.add_paragraph(title)

    items = price_data.get("items", [])
    total_row = price_data.get("total_row")

    # Table: S. No | Component | Price
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "S. No"
    hdr[1].text = "Component"
    hdr[2].text = "Price"

    for item in items:
        row_cells = table.add_row().cells
        row_cells[0].text = str(item.get("s_no", ""))
        row_cells[1].text = str(item.get("label", ""))
        row_cells[2].text = str(item.get("price", ""))

    # Total row
    if total_row:
        row_cells = table.add_row().cells
        row_cells[0].text = ""
        row_cells[1].text = str(total_row.get("label", "Total"))
        row_cells[2].text = str(total_row.get("price", ""))

    # Optional BCA discount row
    if apply_bca and total_row:
        final_total_str = apply_bca_discount_to_price_data(price_data, discount_percent)
        if final_total_str:
            row_cells = table.add_row().cells
            row_cells[0].text = ""
            row_cells[1].text = "Final Total (after 4.5% BCA Discount)"
            row_cells[2].text = final_total_str

    # --- Payment Terms section ---
    if payment_terms:
        doc.add_paragraph("")  # spacing
        doc.add_heading("Payment Terms", level=2)

        pt_table = doc.add_table(rows=1, cols=2)
        pt_hdr = pt_table.rows[0].cells
        pt_hdr[0].text = "Payment Percentage"
        pt_hdr[1].text = "Stage"

        for row in payment_terms:
            perc = str(row.get("Payment Percentage", "")).strip()
            stage = str(row.get("Stage", "")).strip()
            if not perc and not stage:
                continue
            r = pt_table.add_row().cells
            r[0].text = perc
            r[1].text = stage

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ---------- Streamlit UI ----------

uploaded_file = st.file_uploader("Upload costing Excel", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        # Read Overall Costing as DataFrame and convert to CSV for the model
        df = pd.read_excel(uploaded_file, sheet_name="Overall Costing", header=None)
        sheet_csv = df.to_csv(index=False)

        st.subheader("Raw 'Overall Costing' (preview)")
        st.dataframe(df.head(30), use_container_width=True)

        # Button to ask Groq
        if st.button("Ask Groq to build Price Sheet"):
            with st.spinner("Calling Groq to extract high-level Price Sheet..."):
                price_data = call_groq_for_price_sheet(sheet_csv)
            st.session_state["price_data"] = price_data
            st.success("Groq response received and stored.")

        # If we already have price_data (from this or previous click), show editor + payment terms + download
        if "price_data" in st.session_state:
            price_data = st.session_state["price_data"]

            items = price_data.get("items", [])
            if not items:
                st.error("No items returned from Groq. Check prompt or sheet content.")
            else:
                st.subheader("Editable Price Table")

                items_df = pd.DataFrame(items)
                edited_items_df = st.data_editor(
                    items_df,
                    num_rows="dynamic",
                    use_container_width=True,
                    key="price_items_editor",
                )

                # Push edited rows back to dict
                price_data["items"] = edited_items_df.to_dict(orient="records")
                st.session_state["price_data"] = price_data

                st.markdown("---")
                st.subheader("Business Cooperation Agreement (BCA) Discount")

                apply_bca = st.checkbox(
                    "Apply Business Cooperation Agreement Discount (4.5% on Total)",
                    value=False,
                    key="apply_bca_discount",
                )

                # -------- Payment Terms --------
                st.markdown("---")
                st.subheader("Payment Terms")

                # Default payment terms (your last example)
                default_payment_terms = [
                    {"Payment Percentage": "20%", "Stage": "Advance along with LOI/ PO"},
                    {"Payment Percentage": "20%", "Stage": "After DAP Completion"},
                    {"Payment Percentage": "40%", "Stage": "Before Dispatch"},
                    {"Payment Percentage": "10%", "Stage": "Against Installation"},
                    {"Payment Percentage": "10%", "Stage": "Against Handover"},
                ]

                if "payment_terms" not in st.session_state:
                    st.session_state["payment_terms"] = default_payment_terms

                pt_df = pd.DataFrame(st.session_state["payment_terms"])
                edited_pt_df = st.data_editor(
                    pt_df,
                    num_rows="dynamic",
                    use_container_width=True,
                    key="payment_terms_editor",
                )
                st.session_state["payment_terms"] = edited_pt_df.to_dict(orient="records")

                st.markdown("---")
                st.subheader("Download DOCX")

                # Always compute buffer on each run; download_button triggers download
                docx_buffer = build_price_sheet_docx(
                    price_data,
                    apply_bca,
                    payment_terms=st.session_state["payment_terms"],
                    discount_percent=4.5,
                )
                st.download_button(
                    label="Generate & Download Price Sheet (.docx)",
                    data=docx_buffer,
                    file_name="Price_Sheet_Summary_With_Payment_Terms.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )

    except Exception as e:
        st.error(f"Error processing file or calling Groq: {e}")
else:
    st.info("Please upload a costing Excel file to start.")
