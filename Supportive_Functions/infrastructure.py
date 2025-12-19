import streamlit as st
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Infrastructure Section Builder", page_icon="🏗️")

st.title("Infrastructure – Section Builder")

st.write(
    "This will generate the fixed 'Infrastructure' section as a DOCX file, "
    "ready to be merged into the proposal."
)


def build_infrastructure_docx() -> bytes:
    doc = Document()

    # Main section heading
    doc.add_heading("Infrastructure", level=2)

    # a. Fire Protection
    p = doc.add_paragraph()
    p.add_run("a. Fire Protection-").bold = True

    doc.add_paragraph(
        "Falcon’s scope does not cover the design or provision of fire protection infrastructure, "
        "utilities, or related services. It is expected that the customer’s sprinkler contractor "
        "will design and supply the in-rack sprinkler systems, including connectors and mounting "
        "brackets. These designs should be submitted to Falcon for review during the engineering phase."
    )
    doc.add_paragraph(
        "Falcon will work in coordination with the chosen sprinkler supplier to determine the "
        "appropriate locations for the sprinklers and brackets. Modifications may be required to "
        "meet local fire safety regulations, which could influence storage capacity, timelines, and "
        "costs. Any additional sprinkler systems that might affect the overall design must also "
        "undergo review."
    )
    doc.add_paragraph("")

    # b. Power Supply
    p = doc.add_paragraph()
    p.add_run("b. Power Supply-").bold = True

    doc.add_paragraph(
        "The Customer must provide temporary power for installation and permanent power for "
        "commissioning. Protected multi-gang power points for workstations and peripherals will be "
        "supplied by the Customer, with planning for their locations done with the operations and "
        "IT teams."
    )
    doc.add_paragraph("")

    # c. Floor Requirements
    p = doc.add_paragraph()
    p.add_run("c. Floor Requirements-").bold = True

    doc.add_paragraph(
        "The Customer must provide flooring with appropriate loading strength and space at the site. "
        "Falcon assumes that the floor slab will not contain corrosive materials that could affect "
        "standard fixings."
    )
    doc.add_paragraph("")

    # d. Estimated Floor Load
    p = doc.add_paragraph()
    p.add_run("d. Estimated Floor Load-").bold = True

    doc.add_paragraph(
        "Estimated floor loads, including distributed and point loads, will be provided during the "
        "detailed engineering phase of the project."
    )
    doc.add_paragraph("")

    # e. Staging, Laydown and Assembly Area
    p = doc.add_paragraph()
    p.add_run("e. Staging, Laydown and Assembly Area-").bold = True

    doc.add_paragraph(
        "The Customer is required to provide sufficient space on the same floor, adjacent to the "
        "installation site, for staging, storage, and equipment assembly. If such space is "
        "unavailable, offsite locations or areas on different floors may be utilized; however, any "
        "related costs, including additional handling or transportation, will be the Customer’s "
        "responsibility. This may also result in adjustments to the project timeline. Falcon will "
        "specify the necessary requirements for these areas during the design phase as part of "
        "overall project planning and coordination."
    )
    doc.add_paragraph("")

    # f. Site Access and Unloading
    p = doc.add_paragraph()
    p.add_run("f. Site Access and Unloading-").bold = True

    doc.add_paragraph(
        "The Customer is required to allocate sufficient on-site space for parking and staging "
        "shipping containers to facilitate Falcon’s delivery schedule."
    )
    doc.add_paragraph(
        "Additionally, the Customer is responsible for ensuring proper access to the building for "
        "equipment unloading, including providing enough functional dock levellers on all floors "
        "during installation. It is expected that access to the site and installation areas will "
        "remain unobstructed and available around the clock, as needed, to support project operations."
    )
    doc.add_paragraph("")

    # g. Lighting
    p = doc.add_paragraph()
    p.add_run("g. Lighting-").bold = True

    doc.add_paragraph(
        "All lighting is excluded from Falcon’s scope of supply and must be provided by the Customer "
        "or their contractor. This includes lighting for service areas, operational areas, and beneath "
        "platforms and walkways."
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


if st.button("Generate Infrastructure DOCX"):
    doc_bytes = build_infrastructure_docx()
    st.session_state["infrastructure_doc"] = doc_bytes
    st.success("Infrastructure DOCX generated. You can download it below.")

if "infrastructure_doc" in st.session_state:
    st.download_button(
        label="Download Infrastructure (.docx)",
        data=st.session_state["infrastructure_doc"],
        file_name="Infrastructure_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
