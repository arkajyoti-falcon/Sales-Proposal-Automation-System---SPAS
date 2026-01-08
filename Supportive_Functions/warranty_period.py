import streamlit as st
from io import BytesIO
from docx import Document

st.set_page_config(page_title="Warranty Section Builder", page_icon="🛠")

st.title("Warranty Section Builder")

# =========================
# FIXED CONTENT
# =========================

# Fixed coverage bullets – always included
FIXED_COVERAGE = [
    "24 X 7 Telephonic, Email and Remote Service Support when required.",
    "Regular Software updates and Bug Fixes.",
    "Supply of Mechanical and Electrical components in case of failure (excluding damages as mentioned in the Exclusion Clause).",
]

# Fixed exclusions – always included
FIXED_EXCLUSIONS = [
    "Normal wear and tear.",
    "Consumables.",
    "Faulty articles continued:",
    "Failure to comply with the manufacturer's recommendations (logistics documentation, Technical Information Note, retrofit document) and the rules of the trade.",
    "Negligence or abnormal use of equipment.",
    "Anomalies produced by an environment of use, storage or transport that does not comply with the specifications or recommendations of Falcon: packaging, temperature, hygrometry, sector, insulation, etc.",
    "A defect due to a cause external to the supplies and services of Falcon.",
    "Equipment other than that supplied by Falcon.",
    "Items that can be repaired exclusively by Falcon that have been repaired or attempted repairs other than those carried out by Falcon.",
    "Items that fail due to normal wear and tear of one or more of its components or whose tamper-evident seals (varnish, strip, etc.) have been broken or whose serial numbers have been removed or modified.",
    "Items damaged during transport to Falcon due to the use of unsuitable packaging.",
]

# =========================
# VARIABLE: WARRANTY INTRO
# =========================

st.subheader("Warranty Period – Parameters (Variable Part)")

# Warranty type (label only)
warranty_type = st.selectbox(
    "Warranty Type",
    [
        "Standard warranty",
        "Comprehensive warranty",
    ],
)

# Duration
duration = st.text_input(
    "Warranty duration text",
    value="1 year",
    help="Example: '1 year', '12 months', '3 years', '12 months or 15 months from dispatch, whichever is earlier', etc.",
)

# Start condition
start_condition_choice = st.selectbox(
    "Warranty start condition",
    [
        "from the date of beneficiary use.",
        "from the date of commissioning of the system.",
        "from the date of completion of dispatch of the materials, whichever is earlier.",
        "from the date of beneficiary use, max 30 days after readiness of commissioning.",
        "from the date of official communication of material readiness at Falcon end.",
    ],
)

# Extended warranty flag
include_extended = st.checkbox(
    "Include 'Extended Warranty' option after standard warranty?",
    value=True,
)

extended_text_default = (
    "Post completion of the warranty period, the customer can opt for an Extended Warranty."
)
extended_text = ""
if include_extended:
    extended_text = st.text_input(
        "Extended warranty sentence",
        value=extended_text_default,
    )

# AMC / Hotline flag
include_amc = st.checkbox(
    "Include AMC / Hotline clause?",
    value=False,
)

amc_text_default = (
    "After the warranty year, an additional 2 years are covered under AMC and 24x7 Hotline support."
)
amc_text = ""
if include_amc:
    amc_text = st.text_input(
        "AMC / Hotline sentence",
        value=amc_text_default,
    )

# Transport note
include_transport_note = st.checkbox(
    "Include transportation note for sending faulty items to Falcon?",
    value=False,
)

transport_text_default = (
    "Note: Transportation cost for sending faulty items to Falcon factory from the site will be in the customer’s scope."
)
transport_text = ""
if include_transport_note:
    transport_text = st.text_area(
        "Transportation note",
        value=transport_text_default,
    )

st.markdown("---")


def build_warranty_intro(
    warranty_type: str,
    duration: str,
    start_condition: str,
    extended_text: str,
    amc_text: str,
) -> str:
    """
    Build the main intro paragraph for the warranty section.
    """
    base = f"Falcon’s offered System comes with a {warranty_type.lower()} of {duration} (starts {start_condition})"
    if not base.strip().endswith("."):
        base += "."

    extra_sentences = []
    if extended_text:
        extra_sentences.append(extended_text.strip())
    if amc_text:
        extra_sentences.append(amc_text.strip())

    if extra_sentences:
        base += " " + " ".join(extra_sentences)

    return base


def build_docx(
    warranty_type: str,
    duration: str,
    start_condition: str,
    extended_text: str,
    amc_text: str,
    transport_text: str,
):
    doc = Document()

    # Heading 2: Warranty Period
    doc.add_heading("Warranty Period", level=2)

    # Intro paragraph (variable)
    intro_para = build_warranty_intro(
        warranty_type=warranty_type,
        duration=duration,
        start_condition=start_condition,
        extended_text=extended_text,
        amc_text=amc_text,
    )
    doc.add_paragraph(intro_para)

    # Coverage (fixed)
    doc.add_paragraph("The warranty covers the following support:")
    for item in FIXED_COVERAGE:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    # Exclusions (fixed)
    doc.add_paragraph("The warranty does not apply to the replacement or repair of:")
    for item in FIXED_EXCLUSIONS:
        # For the "Faulty articles continued:" we keep it as a separate bullet
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    # Optional transport note
    if transport_text:
        doc.add_paragraph(transport_text.strip())

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


st.subheader("Generate Warranty DOCX")

if st.button("Generate Warranty Section DOCX"):
    docx_buffer = build_docx(
        warranty_type=warranty_type,
        duration=duration,
        start_condition=start_condition_choice,
        extended_text=extended_text,
        amc_text=amc_text,
        transport_text=transport_text,
    )

    st.download_button(
        label="Download Warranty Section (.docx)",
        data=docx_buffer,
        file_name="Warranty_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
