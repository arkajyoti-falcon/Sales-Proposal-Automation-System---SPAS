import streamlit as st
from io import BytesIO
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

st.set_page_config(page_title="Falcon SCADA Section", page_icon="📊")

st.title("Falcon Visual Inspection System (SCADA) – DOCX Builder")

# ---- Config: backend image paths (adjust as per your repo/structure) ----
IMAGE_PATHS = {
    "field_data": "FIXED_IMAGE\\main_plc.PNG",   # 1. FIELD DATA ACQUISITION
    "vis_1": "FIXED_IMAGE\\animated_sys_v1.PNG",        # 2. ANIMATED SYSTEM VISUALIZATION – view 1
    "vis_2": "FIXED_IMAGE\\animated_sys_v2.PNG",        # 2. ANIMATED SYSTEM VISUALIZATION – view 2
    "alarm": "FIXED_IMAGE\\alarm.PNG",        # 3. ALARM MANAGEMENT
}


def add_centered_image_from_path(doc: Document, path: str, width_in=5.5):
    """
    Helper to add a centered image to a DOCX document from a file path.
    If file doesn't exist, silently skip.
    """
    if not path or not os.path.exists(path):
        return
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_scada_docx(client_name: str) -> BytesIO:
    doc = Document()

    # Heading 2
    doc.add_heading("Falcon’s Visual Inspection System (SCADA)", level=2)

    # Intro text
    doc.add_paragraph(
        "SCADA stands for Supervisory Control and Data Acquisition. It is a system of hardware "
        "and software components that allows for remote monitoring, control, and data acquisition "
        "of industrial processes or facilities."
    )

    doc.add_paragraph(
        f"The Visualization system provided by FALCON (or SCADA) allows the monitoring and control of "
        f"the different systems delivered for the {client_name}. This SCADA system receives from each "
        f"monitored sub-system all information on their operating status in real time."
    )

    doc.add_paragraph("At the system monitoring level, the functions performed are:")

    functions = [
        "Field data acquisition.",
        "Animated visualization of equipment.",
        "Representation of the operating mode of the system (nominal, contingency, etc.).",
        "Alarm management.",
        "Alarm history management.",
        "Diagnostic help.",
        "Failure detection.",
        "Equipment control.",
        "Statistics on equipment operation.",
        "Historical Statistical Report.",
        "Recording and archiving.",
        "Safety operator interface.",
    ]
    for item in functions:
        doc.add_paragraph(item, style="List Bullet")

    # 1. FIELD DATA ACQUISITION
    doc.add_heading("1. FIELD DATA ACQUISITION", level=3)
    doc.add_paragraph(
        "The field data acquisition function is performed by the SCADA system connected to the sorters' PLCs. "
        "The communication with the PLCs is done using equipped CPU cards that are able to manage the "
        "communication with the PLC on the Industrial Ethernet network, without overloading the server."
    )
    doc.add_paragraph(
        "The acquisition of signals from external systems is carried out via the PLC, using dry contacts."
    )

    # Image – Field data acquisition (center)
    add_centered_image_from_path(doc, IMAGE_PATHS.get("field_data"))

    # 2. ANIMATED SYSTEM VISUALIZATION
    doc.add_heading("2. ANIMATED SYSTEM VISUALIZATION", level=3)
    doc.add_paragraph(
        "The animated view represents the dynamic graphical user interface that allows real-time monitoring "
        "of the controlled systems and the execution of their control procedures."
    )
    doc.add_paragraph(
        "The various states of the digital signals are displayed on the screen by graphic symbols. "
        "All views are web based with responsive capabilities when required so that various devices can be "
        "used to display status and information on PC. The types of information that can be displayed on each "
        "type of device will be discussed during the design phase of the project."
    )

    # Images – Visualization screens (centered)
    add_centered_image_from_path(doc, IMAGE_PATHS.get("vis_1"))
    add_centered_image_from_path(doc, IMAGE_PATHS.get("vis_2"))

    # 3. ALARM MANAGEMENT
    doc.add_heading("3. ALARM MANAGEMENT", level=3)
    doc.add_paragraph(
        "The alarm pages display a series of information to identify the nature of the alarm or event, "
        "the elements involved and the time."
    )

    doc.add_paragraph("The alarm management includes:")

    alarm_points = [
        "The Alarm name.",
        "The date on which the above-mentioned alarm was activated.",
        "The moment the alarm goes off.",
        "The date on which the alarm is activated again.",
        "The moment the alarm is activated again.",
        "The sub-system that activated the alarm.",
        "The area where the alarm was triggered.",
        "The name of the beacon that activated the alarm.",
        "The alarm state.",
        "The alarm value.",
        "A complete alarm description.",
    ]
    for item in alarm_points:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph(
        "The historic data of each alarm is stored in the SCADA database. This type of fault detected includes:"
    )

    fault_points = [
        "Sensors.",
        "Drive Fault.",
        "Lack of monitored equipment.",
        "Etc.",
    ]
    for item in fault_points:
        doc.add_paragraph(item, style="List Bullet")

    # Image – Alarm management (center)
    add_centered_image_from_path(doc, IMAGE_PATHS.get("alarm"))

    # Export buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ------------ UI ------------

client_name = st.text_input("Client Name (used inside SCADA text)", value="Zepto")

st.markdown("---")
if st.button("Generate SCADA DOCX"):
    docx_buffer = build_scada_docx(client_name.strip() or "Client")
    st.download_button(
        label="Download 'Falcon SCADA' Section (.docx)",
        data=docx_buffer,
        file_name="Falcon_SCADA_Section.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
